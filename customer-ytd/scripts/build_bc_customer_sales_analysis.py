from __future__ import annotations

import argparse
import math
import os
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any
from collections.abc import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

WORKSPACE = Path(os.environ.get("CUSTOMER_YTD_WORKSPACE", Path.cwd())).resolve()
sys.path.append(str(WORKSPACE / "src"))
sys.path.append(str(WORKSPACE / "scripts"))

from build_bc_weekly_dashboard import (  # noqa: E402
    BASE_CURRENCY,
    BusinessCentralClient,
    BusinessCentralConfig,
    BusinessCentralError,
    fetch_api_rows_by_month,
    fetch_exchange_rates,
    fx_rate,
    last_day,
    normal_text,
    parse_date,
)


OUTPUT_DIR = Path("outputs/bc-customer-sales-analysis")

TITLE_FILL = "193B4D"
SECTION_FILL = "3D6F7E"
ACCENT_FILL = "E8F1F4"
LIGHT_FILL = "F5F8FA"
GOOD_FILL = "E2F0D9"
BAD_FILL = "FCE4D6"
TEXT = "1F2933"
MUTED = "667085"

SWISS_COUNTRIES = {"CH", "CHE", "SWITZERLAND", "SCHWEIZ", "SUISSE"}
REGIONS = ["Swiss", "Samnaun", "Export"]
SAMNAUN_CITIES = {"samnaun", "samnaun dorf", "martina"}


@dataclass(frozen=True)
class Period:
    label: str
    start: date
    end: date


@dataclass(frozen=True)
class SalesRow:
    posting_date: date
    period: str
    doc_no: str
    doc_type: str
    customer_no: str
    customer_name: str
    country: str
    city: str
    region: str
    item_no: str
    description: str
    quantity: float
    sales_chf: float
    currency_code: str
    sales_original: float


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build customer sales analysis from posted BC sales invoices and credit memos."
    )
    parser.add_argument(
        "--as-of",
        default=date.today().isoformat(),
        help="Report date in YYYY-MM-DD format. Defaults to today.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=OUTPUT_DIR,
        help="Folder where the workbook will be saved.",
    )
    parser.add_argument(
        "--top",
        type=int,
        default=20,
        help="Number of top customers to show in dashboard and regional charts.",
    )
    return parser.parse_args()


def safe_replace_year(value: date, year: int) -> date:
    try:
        return value.replace(year=year)
    except ValueError:
        return value.replace(year=year, day=28)


def ytd_period(as_of: date) -> Period:
    return Period("YTD", date(as_of.year, 1, 1), as_of)


def mtd_period(as_of: date) -> Period:
    return Period("MTD", as_of.replace(day=1), as_of)


def prior_year_same_dates(period: Period) -> Period:
    return comparison_same_dates(period, 1)


def comparison_same_dates(period: Period, years_back: int) -> Period:
    return Period(
        f"{period.label} -{years_back}Y",
        safe_replace_year(period.start, period.start.year - years_back),
        safe_replace_year(period.end, period.end.year - years_back),
    )


def in_period(value: date, period: Period) -> bool:
    return period.start <= value <= period.end


def clean_number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    try:
        result = float(value)
    except (TypeError, ValueError):
        return 0.0
    if math.isnan(result) or math.isinf(result):
        return 0.0
    return result


def classify_region(country: str, city: str) -> str:
    if normal_text(city).casefold() in SAMNAUN_CITIES:
        return "Samnaun"
    if normal_text(country).upper() in SWISS_COUNTRIES:
        return "Swiss"
    return "Export"


def fetch_customer_sales_rows(
    client: BusinessCentralClient,
    rates: dict[str, list[tuple[date, float]]],
    current_ytd: Period,
    prior_ytd: Period,
    two_year_ytd: Period,
) -> list[SalesRow]:
    rows: list[SalesRow] = []
    fetch_plan = [
        ("Current", current_ytd),
        ("Prior year", prior_ytd),
        ("2024", two_year_ytd),
    ]
    configs = [
        ("salesInvoices", "salesInvoiceLines", "Invoice", 1),
        ("salesCreditMemos", "salesCreditMemoLines", "Credit memo", -1),
    ]
    for period_label, period in fetch_plan:
        for resource, line_resource, doc_type, sign in configs:
            headers = fetch_api_rows_by_month(
                client,
                resource,
                date_field="postingDate",
                start=period.start,
                end=period.end,
                order_by="postingDate asc",
                expand=line_resource,
            )
            for header in headers:
                if normal_text(header.get("status")).casefold() in {"draft", "canceled", "cancelled"}:
                    continue
                posting_date = parse_date(header["postingDate"])
                currency = normal_text(header.get("currencyCode")).upper()
                rate = fx_rate(rates, currency, posting_date)
                country = normal_text(header.get("sellToCountry"))
                city = normal_text(header.get("sellToCity"))
                region = classify_region(country, city)
                for line in header.get(line_resource) or []:
                    amount_original = clean_number(line.get("amountExcludingTax"))
                    amount_chf = amount_original * rate * sign
                    quantity = clean_number(line.get("quantity")) * sign
                    if abs(amount_chf) < 0.000001 and abs(quantity) < 0.000001:
                        continue
                    rows.append(
                        SalesRow(
                            posting_date=posting_date,
                            period=period_label,
                            doc_no=normal_text(header.get("number")),
                            doc_type=doc_type,
                            customer_no=normal_text(header.get("customerNumber")),
                            customer_name=normal_text(header.get("customerName")) or "(unknown)",
                            country=country,
                            city=city,
                            region=region,
                            item_no=normal_text(line.get("lineObjectNumber")),
                            description=normal_text(line.get("description")),
                            quantity=quantity,
                            sales_chf=amount_chf,
                            currency_code=currency or BASE_CURRENCY,
                            sales_original=amount_original * sign,
                        )
                    )
    return rows


def customer_key(row: SalesRow) -> tuple[str, str, str, str, str]:
    return (
        row.region,
        row.customer_no,
        row.customer_name,
        row.country,
        row.city,
    )


def pct_change(current: float, prior: float) -> float | None:
    if abs(prior) < 0.000001:
        return None
    return (current - prior) / prior


def summarize_customers(
    rows: Iterable[SalesRow],
    as_of: date,
) -> list[dict[str, Any]]:
    current_ytd = ytd_period(as_of)
    prior_year_same_dates(current_ytd)
    comparison_same_dates(current_ytd, 2)
    current_mtd = mtd_period(as_of)
    prior_mtd = prior_year_same_dates(current_mtd)
    two_year_mtd = comparison_same_dates(current_mtd, 2)
    grouped: dict[tuple[str, str, str, str, str], dict[str, Any]] = defaultdict(lambda: defaultdict(float))
    for row in rows:
        key = customer_key(row)
        grouped[key]["quantity"] += row.quantity
        grouped[key]["line_count"] += 1
        grouped[key].setdefault("docs_current", set())
        grouped[key].setdefault("docs_prior", set())
        grouped[key].setdefault("docs_2024", set())
        if row.period == "Current":
            grouped[key]["ytd_current"] += row.sales_chf
            grouped[key]["qty_current"] += row.quantity
            grouped[key]["docs_current"].add((row.doc_type, row.doc_no))
            if in_period(row.posting_date, current_mtd):
                grouped[key]["mtd_current"] += row.sales_chf
        elif row.period == "Prior year":
            grouped[key]["ytd_prior"] += row.sales_chf
            grouped[key]["qty_prior"] += row.quantity
            grouped[key]["docs_prior"].add((row.doc_type, row.doc_no))
            if in_period(row.posting_date, prior_mtd):
                grouped[key]["mtd_prior"] += row.sales_chf
        elif row.period == "2024":
            grouped[key]["ytd_2024"] += row.sales_chf
            grouped[key]["qty_2024"] += row.quantity
            grouped[key]["docs_2024"].add((row.doc_type, row.doc_no))
            if in_period(row.posting_date, two_year_mtd):
                grouped[key]["mtd_2024"] += row.sales_chf

    records: list[dict[str, Any]] = []
    total_current = sum(values.get("ytd_current", 0.0) for values in grouped.values())
    for key, values in grouped.items():
        region, no, name, country, city = key
        current = values.get("ytd_current", 0.0)
        prior = values.get("ytd_prior", 0.0)
        two_year = values.get("ytd_2024", 0.0)
        mtd_current = values.get("mtd_current", 0.0)
        mtd_prior = values.get("mtd_prior", 0.0)
        mtd_two_year = values.get("mtd_2024", 0.0)
        if abs(current) < 0.000001 and abs(prior) < 0.000001 and abs(two_year) < 0.000001:
            continue
        records.append(
            {
                "region": region,
                "customer_no": no,
                "customer_name": name,
                "country": country,
                "city": city,
                "ytd_current": current,
                "ytd_prior": prior,
                "ytd_2024": two_year,
                "ytd_delta": current - prior,
                "ytd_pct": pct_change(current, prior),
                "ytd_delta_2024": current - two_year,
                "ytd_pct_2024": pct_change(current, two_year),
                "mtd_current": mtd_current,
                "mtd_prior": mtd_prior,
                "mtd_2024": mtd_two_year,
                "mtd_delta": mtd_current - mtd_prior,
                "mtd_pct": pct_change(mtd_current, mtd_prior),
                "mtd_delta_2024": mtd_current - mtd_two_year,
                "mtd_pct_2024": pct_change(mtd_current, mtd_two_year),
                "qty_current": values.get("qty_current", 0.0),
                "qty_prior": values.get("qty_prior", 0.0),
                "qty_2024": values.get("qty_2024", 0.0),
                "doc_count_current": len(values.get("docs_current", set())),
                "doc_count_prior": len(values.get("docs_prior", set())),
                "doc_count_2024": len(values.get("docs_2024", set())),
                "share_current": None if abs(total_current) < 0.000001 else current / total_current,
            }
        )
    return sorted(records, key=lambda item: (REGIONS.index(item["region"]), -abs(item["ytd_current"])))


def summarize_regions(customer_records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = defaultdict(lambda: defaultdict(float))
    for rec in customer_records:
        region = rec["region"]
        grouped[region]["ytd_current"] += rec["ytd_current"]
        grouped[region]["ytd_prior"] += rec["ytd_prior"]
        grouped[region]["ytd_2024"] += rec["ytd_2024"]
        grouped[region]["mtd_current"] += rec["mtd_current"]
        grouped[region]["mtd_prior"] += rec["mtd_prior"]
        grouped[region]["mtd_2024"] += rec["mtd_2024"]
        grouped[region]["qty_current"] += rec["qty_current"]
        grouped[region]["qty_prior"] += rec["qty_prior"]
        grouped[region]["qty_2024"] += rec["qty_2024"]
        grouped[region]["doc_count_current"] += rec["doc_count_current"]
        grouped[region]["doc_count_prior"] += rec["doc_count_prior"]
        grouped[region]["doc_count_2024"] += rec["doc_count_2024"]
        if abs(rec["ytd_current"]) >= 0.000001:
            grouped[region]["customer_count_current"] += 1
        if abs(rec["ytd_prior"]) >= 0.000001:
            grouped[region]["customer_count_prior"] += 1
        if abs(rec["ytd_2024"]) >= 0.000001:
            grouped[region]["customer_count_2024"] += 1
    rows = []
    total_current = sum(grouped[region].get("ytd_current", 0.0) for region in REGIONS)
    for region in REGIONS:
        values = grouped[region]
        current = values.get("ytd_current", 0.0)
        prior = values.get("ytd_prior", 0.0)
        two_year = values.get("ytd_2024", 0.0)
        mtd_current = values.get("mtd_current", 0.0)
        mtd_prior = values.get("mtd_prior", 0.0)
        mtd_two_year = values.get("mtd_2024", 0.0)
        rows.append(
            {
                "region": region,
                "ytd_current": current,
                "ytd_prior": prior,
                "ytd_2024": two_year,
                "ytd_delta": current - prior,
                "ytd_pct": pct_change(current, prior),
                "ytd_delta_2024": current - two_year,
                "ytd_pct_2024": pct_change(current, two_year),
                "mtd_current": mtd_current,
                "mtd_prior": mtd_prior,
                "mtd_2024": mtd_two_year,
                "mtd_delta": mtd_current - mtd_prior,
                "mtd_pct": pct_change(mtd_current, mtd_prior),
                "mtd_delta_2024": mtd_current - mtd_two_year,
                "mtd_pct_2024": pct_change(mtd_current, mtd_two_year),
                "qty_current": values.get("qty_current", 0.0),
                "qty_prior": values.get("qty_prior", 0.0),
                "qty_2024": values.get("qty_2024", 0.0),
                "doc_count_current": values.get("doc_count_current", 0.0),
                "doc_count_prior": values.get("doc_count_prior", 0.0),
                "doc_count_2024": values.get("doc_count_2024", 0.0),
                "customer_count_current": values.get("customer_count_current", 0.0),
                "customer_count_prior": values.get("customer_count_prior", 0.0),
                "customer_count_2024": values.get("customer_count_2024", 0.0),
                "share_current": None if abs(total_current) < 0.000001 else current / total_current,
            }
        )
    return rows


def monthly_summary(rows: Iterable[SalesRow], as_of: date) -> list[dict[str, Any]]:
    result = []
    for month in range(1, as_of.month + 1):
        start = date(as_of.year, month, 1)
        end = min(last_day(start), as_of if month == as_of.month else last_day(start))
        current = Period(start.strftime("%b"), start, end)
        prior = prior_year_same_dates(current)
        two_year = comparison_same_dates(current, 2)
        row: dict[str, Any] = {
            "month": start.strftime("%b"),
            "current_total": 0.0,
            "prior_total": 0.0,
            "2024_total": 0.0,
        }
        for region in REGIONS:
            row[f"{region}_current"] = 0.0
            row[f"{region}_prior"] = 0.0
            row[f"{region}_2024"] = 0.0
        for sales in rows:
            if sales.period == "Current" and in_period(sales.posting_date, current):
                row["current_total"] += sales.sales_chf
                row[f"{sales.region}_current"] += sales.sales_chf
            elif sales.period == "Prior year" and in_period(sales.posting_date, prior):
                row["prior_total"] += sales.sales_chf
                row[f"{sales.region}_prior"] += sales.sales_chf
            elif sales.period == "2024" and in_period(sales.posting_date, two_year):
                row["2024_total"] += sales.sales_chf
                row[f"{sales.region}_2024"] += sales.sales_chf
        result.append(row)
    return result


def build_workbook(
    *,
    as_of: date,
    rows: list[SalesRow],
    output_path: Path,
    top: int,
) -> None:
    current_ytd = ytd_period(as_of)
    prior_ytd = prior_year_same_dates(current_ytd)
    two_year_ytd = comparison_same_dates(current_ytd, 2)
    current_mtd = mtd_period(as_of)
    prior_mtd = prior_year_same_dates(current_mtd)
    two_year_mtd = comparison_same_dates(current_mtd, 2)
    customer_records = summarize_customers(rows, as_of)
    region_records = summarize_regions(customer_records)
    month_records = monthly_summary(rows, as_of)

    wb = Workbook()
    wb.remove(wb.active)
    dashboard = wb.create_sheet("Dashboard")
    regional = wb.create_sheet("Region Summary")
    customers = wb.create_sheet("Customer Detail")
    trend = wb.create_sheet("Monthly Trend")
    for region in REGIONS:
        wb.create_sheet(f"{region} Customers")
    raw = wb.create_sheet("Source Lines")
    notes = wb.create_sheet("Source Notes")

    build_dashboard(dashboard, as_of, current_ytd, prior_ytd, two_year_ytd, region_records, customer_records, month_records, top)
    build_region_summary(regional, as_of, region_records)
    build_customer_detail(customers, customer_records)
    build_monthly_trend(trend, month_records)
    for region in REGIONS:
        build_region_customer_sheet(wb[f"{region} Customers"], region, customer_records, top)
    build_source_lines(raw, rows)
    build_source_notes(notes, as_of, current_ytd, prior_ytd, two_year_ytd, current_mtd, prior_mtd, two_year_mtd)

    for ws in wb.worksheets:
        polish_sheet(ws)
    wb["Dashboard"].freeze_panes = None
    wb["Source Lines"].freeze_panes = "A4"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def add_title(ws: Worksheet, title: str, subtitle: str, merge_to: str = "H") -> None:
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=20, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=TITLE_FILL)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Aptos", size=10, color="FFFFFF")
    ws["A2"].fill = PatternFill("solid", fgColor=TITLE_FILL)
    ws.merge_cells(f"A1:{merge_to}1")
    ws.merge_cells(f"A2:{merge_to}2")


def add_section(ws: Worksheet, row: int, col: int, label: str, width: int = 1) -> None:
    cell = ws.cell(row, col, label)
    cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=SECTION_FILL)
    if width > 1:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + width - 1)


def build_dashboard(
    ws: Worksheet,
    as_of: date,
    current_ytd: Period,
    prior_ytd: Period,
    two_year_ytd: Period,
    region_records: list[dict[str, Any]],
    customer_records: list[dict[str, Any]],
    month_records: list[dict[str, Any]],
    top: int,
) -> None:
    add_title(
        ws,
        "Sales per Customer vs 2025 and 2024",
        f"Business Central posted sales invoices and posted credit notes | Reporting through {as_of:%Y-%m-%d}",
        "P",
    )

    total_current = sum(r["ytd_current"] for r in region_records)
    total_prior = sum(r["ytd_prior"] for r in region_records)
    total_2024 = sum(r["ytd_2024"] for r in region_records)
    total_mtd = sum(r["mtd_current"] for r in region_records)
    total_mtd_prior = sum(r["mtd_prior"] for r in region_records)
    total_mtd_2024 = sum(r["mtd_2024"] for r in region_records)
    kpis = [
        ("YTD sales", total_current, total_prior, total_2024, total_current - total_prior, total_current - total_2024, pct_change(total_current, total_prior)),
        ("MTD sales", total_mtd, total_mtd_prior, total_mtd_2024, total_mtd - total_mtd_prior, total_mtd - total_mtd_2024, pct_change(total_mtd, total_mtd_prior)),
        ("Active customers", sum(r["customer_count_current"] for r in region_records), sum(r["customer_count_prior"] for r in region_records), sum(r["customer_count_2024"] for r in region_records), None, None, None),
        ("YTD period", f"{current_ytd.start:%Y-%m-%d} to {current_ytd.end:%Y-%m-%d}", f"{prior_ytd.start:%Y-%m-%d} to {prior_ytd.end:%Y-%m-%d}", f"{two_year_ytd.start:%Y-%m-%d} to {two_year_ytd.end:%Y-%m-%d}", None, None, None),
    ]
    ws["A4"] = "Executive summary"
    add_section(ws, 4, 1, "Executive summary", 7)
    ws.append(["Metric", "2026", "2025", "2024", "Change vs 2025", "Change vs 2024", "Change % vs 2025"])
    for row in kpis:
        ws.append(list(row))
    style_table(ws, "A5:G9", "DashboardKPIs")

    region_headers = [
        "Region",
        "YTD 2026 CHF",
        "YTD 2025 CHF",
        "YTD 2024 CHF",
        "Change vs 2025 CHF",
        "Change vs 2024 CHF",
        "Change % vs 2025",
        "Share",
        "Customers",
    ]
    ws["A12"] = "Region comparison"
    add_section(ws, 12, 1, "Region comparison", len(region_headers))
    ws.append(region_headers)
    for rec in region_records:
        ws.append(
            [
                rec["region"],
                rec["ytd_current"],
                rec["ytd_prior"],
                rec["ytd_2024"],
                rec["ytd_delta"],
                rec["ytd_delta_2024"],
                rec["ytd_pct"],
                rec["share_current"],
                rec["customer_count_current"],
            ]
        )
    style_table(ws, "A13:I16", "DashboardRegions")

    top_records = sorted(customer_records, key=lambda r: abs(r["ytd_delta"]), reverse=True)[:top]
    start = 19
    add_section(ws, start, 1, f"Top {len(top_records)} customer movements by absolute YTD change vs 2025", 10)
    ws.append(["Customer", "Region", "YTD 2026 CHF", "YTD 2025 CHF", "YTD 2024 CHF", "Change vs 2025 CHF", "Change % vs 2025", "Change vs 2024 CHF", "City", "Country"])
    for rec in top_records:
        ws.append(
            [
                rec["customer_name"],
                rec["region"],
                rec["ytd_current"],
                rec["ytd_prior"],
                rec["ytd_2024"],
                rec["ytd_delta"],
                rec["ytd_pct"],
                rec["ytd_delta_2024"],
                rec["city"],
                rec["country"],
            ]
        )
    if top_records:
        style_table(ws, f"A{start + 1}:J{start + len(top_records) + 1}", "DashboardCustomerMovements")

    chart_start = 45
    add_section(ws, chart_start, 1, "Chart data", 7)
    ws.append(["Month", "2026 YTD run-rate", "2025 same dates", "2024 same dates"])
    running_current = 0.0
    running_prior = 0.0
    running_2024 = 0.0
    for rec in month_records:
        running_current += rec["current_total"]
        running_prior += rec["prior_total"]
        running_2024 += rec["2024_total"]
        ws.append([rec["month"], running_current, running_prior, running_2024])
    style_table(ws, f"A{chart_start + 1}:D{chart_start + len(month_records) + 1}", "DashboardTrendData")

    add_dashboard_charts(ws, len(region_records), len(top_records), len(month_records), chart_start)


def build_region_summary(ws: Worksheet, as_of: date, region_records: list[dict[str, Any]]) -> None:
    add_title(ws, "Region Summary", f"Swiss excludes Samnaun area; Samnaun is city Samnaun, Samnaun Dorf, or Martina | {as_of:%Y-%m-%d}", "P")
    headers = [
        "Region",
        "YTD 2026 CHF",
        "YTD 2025 CHF",
        "YTD 2024 CHF",
        "YTD Change vs 2025 CHF",
        "YTD Change vs 2024 CHF",
        "YTD Change % vs 2025",
        "YTD Change % vs 2024",
        "MTD 2026 CHF",
        "MTD 2025 CHF",
        "MTD 2024 CHF",
        "MTD Change vs 2025 CHF",
        "MTD Change vs 2024 CHF",
        "MTD Change % vs 2025",
        "MTD Change % vs 2024",
        "Qty Current",
        "Qty 2025",
        "Qty 2024",
        "Documents Current",
        "Documents 2025",
        "Documents 2024",
        "Customers Current",
        "Customers 2025",
        "Customers 2024",
        "YTD Share",
    ]
    ws.append([])
    ws.append(headers)
    for rec in region_records:
        ws.append([rec.get(header_key(header)) for header in headers])
    style_table(ws, f"A4:Y{ws.max_row}", "RegionSummary")


def build_customer_detail(ws: Worksheet, records: list[dict[str, Any]]) -> None:
    add_title(ws, "Customer Detail", "All customers with 2026, 2025, or 2024 sales, grouped by requested region", "W")
    headers = [
        "Region",
        "Customer No.",
        "Customer",
        "Country",
        "City",
        "YTD 2026 CHF",
        "YTD 2025 CHF",
        "YTD 2024 CHF",
        "YTD Change vs 2025 CHF",
        "YTD Change % vs 2025",
        "YTD Change vs 2024 CHF",
        "YTD Change % vs 2024",
        "MTD 2026 CHF",
        "MTD 2025 CHF",
        "MTD 2024 CHF",
        "MTD Change vs 2025 CHF",
        "MTD Change % vs 2025",
        "MTD Change vs 2024 CHF",
        "MTD Change % vs 2024",
        "Qty Current",
        "Qty 2025",
        "Qty 2024",
        "Documents Current",
        "Documents 2025",
        "Documents 2024",
        "YTD Share",
    ]
    ws.append([])
    ws.append(headers)
    for rec in records:
        ws.append(
            [
                rec["region"],
                rec["customer_no"],
                rec["customer_name"],
                rec["country"],
                rec["city"],
                rec["ytd_current"],
                rec["ytd_prior"],
                rec["ytd_2024"],
                rec["ytd_delta"],
                rec["ytd_pct"],
                rec["ytd_delta_2024"],
                rec["ytd_pct_2024"],
                rec["mtd_current"],
                rec["mtd_prior"],
                rec["mtd_2024"],
                rec["mtd_delta"],
                rec["mtd_pct"],
                rec["mtd_delta_2024"],
                rec["mtd_pct_2024"],
                rec["qty_current"],
                rec["qty_prior"],
                rec["qty_2024"],
                rec["doc_count_current"],
                rec["doc_count_prior"],
                rec["doc_count_2024"],
                rec["share_current"],
            ]
        )
    if ws.max_row >= 4:
        style_table(ws, f"A4:Y{ws.max_row}", "CustomerDetail")
        add_delta_formatting(ws, f"I5:I{ws.max_row}")
        add_delta_formatting(ws, f"K5:K{ws.max_row}")
        add_delta_formatting(ws, f"P5:P{ws.max_row}")
        add_delta_formatting(ws, f"R5:R{ws.max_row}")


def build_monthly_trend(ws: Worksheet, records: list[dict[str, Any]]) -> None:
    add_title(ws, "Monthly Trend", "2026, 2025, and 2024 sales by month and region", "R")
    headers = [
        "Month",
        "2026 Total CHF",
        "2025 Total CHF",
        "2024 Total CHF",
        "Swiss Current CHF",
        "Swiss 2025 CHF",
        "Swiss 2024 CHF",
        "Samnaun Current CHF",
        "Samnaun 2025 CHF",
        "Samnaun 2024 CHF",
        "Export Current CHF",
        "Export 2025 CHF",
        "Export 2024 CHF",
    ]
    ws.append([])
    ws.append(headers)
    for rec in records:
        ws.append(
            [
                rec["month"],
                rec["current_total"],
                rec["prior_total"],
                rec["2024_total"],
                rec["Swiss_current"],
                rec["Swiss_prior"],
                rec["Swiss_2024"],
                rec["Samnaun_current"],
                rec["Samnaun_prior"],
                rec["Samnaun_2024"],
                rec["Export_current"],
                rec["Export_prior"],
                rec["Export_2024"],
            ]
        )
    style_table(ws, f"A4:M{ws.max_row}", "MonthlyTrend")

    chart = LineChart()
    chart.title = "Monthly sales by region and year"
    chart.y_axis.title = "CHF"
    chart.x_axis.title = "Month"
    data = Reference(ws, min_col=5, max_col=13, min_row=4, max_row=ws.max_row)
    cats = Reference(ws, min_col=1, min_row=5, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    set_series_titles(
        chart,
        [
            "Swiss 2026",
            "Swiss 2025",
            "Swiss 2024",
            "Samnaun 2026",
            "Samnaun 2025",
            "Samnaun 2024",
            "Export 2026",
            "Export 2025",
            "Export 2024",
        ],
    )
    chart.height = 11
    chart.width = 24
    ws.add_chart(chart, "K4")


def build_region_customer_sheet(ws: Worksheet, region: str, records: list[dict[str, Any]], top: int) -> None:
    region_records = [rec for rec in records if rec["region"] == region]
    add_title(ws, f"{region} Customers", "Customer-level YTD and MTD comparison for 2026, 2025, and 2024", "Q")
    headers = [
        "Customer No.",
        "Customer",
        "Country",
        "City",
        "YTD 2026 CHF",
        "YTD 2025 CHF",
        "YTD 2024 CHF",
        "YTD Change vs 2025 CHF",
        "YTD Change % vs 2025",
        "YTD Change vs 2024 CHF",
        "YTD Change % vs 2024",
        "MTD 2026 CHF",
        "MTD 2025 CHF",
        "MTD 2024 CHF",
        "MTD Change vs 2025 CHF",
        "MTD Change vs 2024 CHF",
        "YTD Share",
    ]
    ws.append([])
    ws.append(headers)
    for rec in region_records:
        ws.append(
            [
                rec["customer_no"],
                rec["customer_name"],
                rec["country"],
                rec["city"],
                rec["ytd_current"],
                rec["ytd_prior"],
                rec["ytd_2024"],
                rec["ytd_delta"],
                rec["ytd_pct"],
                rec["ytd_delta_2024"],
                rec["ytd_pct_2024"],
                rec["mtd_current"],
                rec["mtd_prior"],
                rec["mtd_2024"],
                rec["mtd_delta"],
                rec["mtd_delta_2024"],
                rec["share_current"],
            ]
        )
    if region_records:
        style_table(ws, f"A4:Q{ws.max_row}", f"{region}Customers".replace(" ", ""))
        add_delta_formatting(ws, f"H5:H{ws.max_row}")
        add_delta_formatting(ws, f"J5:J{ws.max_row}")
        add_delta_formatting(ws, f"O5:O{ws.max_row}")
        add_delta_formatting(ws, f"P5:P{ws.max_row}")
        chart_rows = min(top, len(region_records))
        chart = BarChart()
        chart.title = f"Top {chart_rows} {region} customers YTD"
        chart.y_axis.title = "CHF"
        data = Reference(ws, min_col=5, max_col=7, min_row=4, max_row=4 + chart_rows)
        cats = Reference(ws, min_col=2, min_row=5, max_row=4 + chart_rows)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        set_series_titles(chart, ["2026", "2025", "2024"])
        chart.height = 10
        chart.width = 22
        ws.add_chart(chart, "O4")


def build_source_lines(ws: Worksheet, rows: list[SalesRow]) -> None:
    add_title(ws, "Source Lines", "Line-level BC source used for the customer analysis", "O")
    headers = [
        "Posting Date",
        "Period",
        "Document Type",
        "Document No.",
        "Region",
        "Customer No.",
        "Customer",
        "Country",
        "City",
        "Item/Line No.",
        "Description",
        "Quantity",
        "Sales CHF",
        "Currency",
        "Sales Original Currency",
    ]
    ws.append([])
    ws.append(headers)
    for row in sorted(rows, key=lambda item: (item.posting_date, item.customer_name, item.doc_no)):
        ws.append(
            [
                row.posting_date,
                row.period,
                row.doc_type,
                row.doc_no,
                row.region,
                row.customer_no,
                row.customer_name,
                row.country,
                row.city,
                row.item_no,
                row.description,
                row.quantity,
                row.sales_chf,
                row.currency_code,
                row.sales_original,
            ]
        )
    if ws.max_row >= 4:
        style_table(ws, f"A4:O{ws.max_row}", "SourceLines")


def build_source_notes(
    ws: Worksheet,
    as_of: date,
    current_ytd: Period,
    prior_ytd: Period,
    two_year_ytd: Period,
    current_mtd: Period,
    prior_mtd: Period,
    two_year_mtd: Period,
) -> None:
    add_title(ws, "Source Notes", "Business Central sources and calculation rules", "C")
    rows = [
        ("Report date", as_of.isoformat()),
        ("Current YTD", f"{current_ytd.start:%Y-%m-%d} to {current_ytd.end:%Y-%m-%d}"),
        ("2025 YTD", f"{prior_ytd.start:%Y-%m-%d} to {prior_ytd.end:%Y-%m-%d}"),
        ("2024 YTD", f"{two_year_ytd.start:%Y-%m-%d} to {two_year_ytd.end:%Y-%m-%d}"),
        ("Current MTD", f"{current_mtd.start:%Y-%m-%d} to {current_mtd.end:%Y-%m-%d}"),
        ("2025 MTD", f"{prior_mtd.start:%Y-%m-%d} to {prior_mtd.end:%Y-%m-%d}"),
        ("2024 MTD", f"{two_year_mtd.start:%Y-%m-%d} to {two_year_mtd.end:%Y-%m-%d}"),
        ("BC sources", "salesInvoices + salesInvoiceLines; salesCreditMemos + salesCreditMemoLines"),
        ("Amount basis", "Line amountExcludingTax, converted to CHF with Business Central currency exchange rates"),
        ("Credit notes", "Posted sales credit memos are included as negative sales and quantity"),
        ("Swiss customers", "sellToCountry is CH/CHE/Switzerland/Schweiz/Suisse and sellToCity is not Samnaun"),
        ("Samnaun customers", "sellToCity equals Samnaun, Samnaun Dorf, or Martina"),
        ("Export customers", "All customers not classified as Swiss or Samnaun"),
    ]
    ws.append([])
    ws.append(["Item", "Value"])
    for row in rows:
        ws.append(list(row))
    style_table(ws, f"A4:B{ws.max_row}", "SourceNotes")


def header_key(header: str) -> str:
    mapping = {
        "Region": "region",
        "YTD 2026 CHF": "ytd_current",
        "YTD 2025 CHF": "ytd_prior",
        "YTD 2024 CHF": "ytd_2024",
        "YTD Change vs 2025 CHF": "ytd_delta",
        "YTD Change vs 2024 CHF": "ytd_delta_2024",
        "YTD Change % vs 2025": "ytd_pct",
        "YTD Change % vs 2024": "ytd_pct_2024",
        "MTD 2026 CHF": "mtd_current",
        "MTD 2025 CHF": "mtd_prior",
        "MTD 2024 CHF": "mtd_2024",
        "MTD Change vs 2025 CHF": "mtd_delta",
        "MTD Change vs 2024 CHF": "mtd_delta_2024",
        "MTD Change % vs 2025": "mtd_pct",
        "MTD Change % vs 2024": "mtd_pct_2024",
        "Qty Current": "qty_current",
        "Qty 2025": "qty_prior",
        "Qty 2024": "qty_2024",
        "Documents Current": "doc_count_current",
        "Documents 2025": "doc_count_prior",
        "Documents 2024": "doc_count_2024",
        "Customers Current": "customer_count_current",
        "Customers 2025": "customer_count_prior",
        "Customers 2024": "customer_count_2024",
        "YTD Share": "share_current",
    }
    return mapping[header]


def add_dashboard_charts(
    ws: Worksheet,
    region_count: int,
    customer_count: int,
    month_count: int,
    chart_start: int,
) -> None:
    bar = BarChart()
    bar.title = "YTD sales by region"
    bar.y_axis.title = "CHF"
    data = Reference(ws, min_col=2, max_col=4, min_row=13, max_row=13 + region_count)
    cats = Reference(ws, min_col=1, min_row=14, max_row=13 + region_count)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    set_series_titles(bar, ["2026", "2025", "2024"])
    bar.height = 8
    bar.width = 14
    ws.add_chart(bar, "J4")

    mix = DoughnutChart()
    mix.title = "Current YTD regional mix"
    mix.holeSize = 55
    data = Reference(ws, min_col=2, min_row=13, max_row=13 + region_count)
    cats = Reference(ws, min_col=1, min_row=14, max_row=13 + region_count)
    mix.add_data(data, titles_from_data=True)
    mix.set_categories(cats)
    mix.dataLabels = DataLabelList()
    mix.dataLabels.showPercent = True
    mix.height = 8
    mix.width = 10
    ws.add_chart(mix, "J20")

    if customer_count:
        plotted_customers = min(customer_count, 12)
        customer = BarChart()
        customer.type = "bar"
        customer.title = "Top customer YTD change"
        customer.y_axis.title = "CHF"
        data = Reference(ws, min_col=6, min_row=20, max_row=20 + plotted_customers)
        cats = Reference(ws, min_col=1, min_row=21, max_row=20 + plotted_customers)
        customer.add_data(data, titles_from_data=True)
        customer.set_categories(cats)
        set_series_titles(customer, ["YTD change"])
        customer.height = 12
        customer.width = 18
        ws.add_chart(customer, "J33")

    trend = LineChart()
    trend.title = "YTD cumulative sales"
    trend.y_axis.title = "CHF"
    trend.x_axis.title = "Month"
    data = Reference(ws, min_col=2, max_col=4, min_row=chart_start + 1, max_row=chart_start + month_count + 1)
    cats = Reference(ws, min_col=1, min_row=chart_start + 2, max_row=chart_start + month_count + 1)
    trend.add_data(data, titles_from_data=True)
    trend.set_categories(cats)
    set_series_titles(trend, ["2026", "2025", "2024"])
    trend.height = 9
    trend.width = 18
    ws.add_chart(trend, "J58")


def set_series_titles(chart: BarChart | LineChart | DoughnutChart, labels: list[str]) -> None:
    for series, label in zip(chart.series, labels):
        series.tx = SeriesLabel(v=label)


def style_table(ws: Worksheet, ref: str, name: str) -> None:
    table = Table(displayName=name[:250], ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def add_delta_formatting(ws: Worksheet, ref: str) -> None:
    ws.conditional_formatting.add(
        ref,
        CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=GOOD_FILL)),
    )
    ws.conditional_formatting.add(
        ref,
        CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=BAD_FILL)),
    )


def polish_sheet(ws: Worksheet) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    thin = Side(style="thin", color="D9E2E6")
    table_headers: list[tuple[int, int, int, int]] = [
        range_boundaries(table.ref) for table in ws.tables.values()
    ]
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(
                name="Aptos",
                size=cell.font.sz or 11,
                bold=cell.font.b,
                italic=cell.font.i,
                color=cell.font.color if cell.font.color and cell.font.color.type == "rgb" else TEXT,
            )
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool):
                header = header_for_cell(ws, cell.row, cell.column, table_headers)
                if "%" in header or "Share" in header:
                    cell.number_format = "0.0%"
                elif "Customer" in header or "Documents" in header:
                    cell.number_format = "#,##0"
                elif "Qty" in header or "Quantity" in header or "CHF" in header or header in {"Current", "Last year", "2024", "2025", "2026"}:
                    cell.number_format = '#,##0.00;[Red]-#,##0.00'
    for column_cells in ws.columns:
        values = [normal_text(cell.value) for cell in column_cells[: min(ws.max_row, 80)]]
        width = min(max(max((len(v) for v in values), default=8) + 2, 10), 42)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 20
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 22
    for cell in ws[1] + ws[2]:
        cell.fill = PatternFill("solid", fgColor=TITLE_FILL)
        cell.font = Font(name="Aptos", bold=cell.row == 1, size=20 if cell.row == 1 else 10, color="FFFFFF")
    for row in ws.iter_rows(min_row=4, max_row=4):
        for cell in row:
            if cell.value:
                cell.font = Font(name="Aptos", bold=True, color=TEXT)
    for col in ws.iter_cols():
        for cell in col:
            if isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"


def header_for_cell(
    ws: Worksheet,
    row: int,
    col: int,
    table_bounds: list[tuple[int, int, int, int]],
) -> str:
    for min_col, min_row, max_col, max_row in table_bounds:
        if min_row < row <= max_row and min_col <= col <= max_col:
            return normal_text(ws.cell(min_row, col).value)
    return normal_text(ws.cell(4, col).value)


def verify_workbook(path: Path) -> None:
    wb = load_workbook(path, data_only=False)
    required = {
        "Dashboard",
        "Region Summary",
        "Customer Detail",
        "Monthly Trend",
        "Swiss Customers",
        "Samnaun Customers",
        "Export Customers",
        "Source Lines",
        "Source Notes",
    }
    missing = required.difference(wb.sheetnames)
    if missing:
        raise RuntimeError(f"Workbook is missing sheets: {', '.join(sorted(missing))}")
    if not wb["Dashboard"]._charts:
        raise RuntimeError("Dashboard has no charts")
    for sheet_name in ["Dashboard", "Region Summary", "Customer Detail", "Monthly Trend"]:
        ws = wb[sheet_name]
        if ws.max_row < 4 or ws.max_column < 2:
            raise RuntimeError(f"Sheet {sheet_name} looks empty")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(
                    err in cell.value for err in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
                ):
                    raise RuntimeError(f"Formula/display error in {ws.title}!{cell.coordinate}: {cell.value}")


def main() -> int:
    args = parse_args()
    as_of = parse_date(args.as_of)
    current_ytd = ytd_period(as_of)
    prior_ytd = prior_year_same_dates(current_ytd)
    two_year_ytd = comparison_same_dates(current_ytd, 2)

    print("Connecting to Business Central...", flush=True)
    client = BusinessCentralClient(BusinessCentralConfig.from_env())

    print("Fetching exchange rates...", flush=True)
    rates = fetch_exchange_rates(client)

    print(f"Fetching posted sales invoices and credit notes for {current_ytd.start:%Y-%m-%d} to {current_ytd.end:%Y-%m-%d}...", flush=True)
    print(f"Fetching comparison period {prior_ytd.start:%Y-%m-%d} to {prior_ytd.end:%Y-%m-%d}...", flush=True)
    print(f"Fetching 2024 comparison period {two_year_ytd.start:%Y-%m-%d} to {two_year_ytd.end:%Y-%m-%d}...", flush=True)
    rows = fetch_customer_sales_rows(client, rates, current_ytd, prior_ytd, two_year_ytd)

    filename = f"Silver_Spirits_Customer_Sales_Analysis_{as_of:%Y-%m-%d}.xlsx"
    output_path = args.output_dir / filename
    print("Building Excel workbook...", flush=True)
    build_workbook(as_of=as_of, rows=rows, output_path=output_path, top=args.top)
    verify_workbook(output_path)

    customer_records = summarize_customers(rows, as_of)
    region_records = summarize_regions(customer_records)
    print(f"Created {output_path}", flush=True)
    print(f"Source lines used: {len(rows):,}", flush=True)
    for rec in region_records:
        print(
            f"{rec['region']}: {rec['ytd_current']:,.2f} CHF vs "
            f"{rec['ytd_prior']:,.2f} CHF (2025), {rec['ytd_2024']:,.2f} CHF (2024)",
            flush=True,
        )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except BusinessCentralError as exc:
        raise SystemExit(str(exc))
