from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill


BLUE_FILL = PatternFill(fill_type="solid", fgColor="FF0070C0")


def normalize_header(value) -> str:
    return str(value or "").strip().casefold()


def value_from_row(values, headers: dict[str, int], *names: str):
    for name in names:
        idx = headers.get(normalize_header(name))
        if idx is not None:
            return values[idx]
    raise KeyError(f"Missing required Stock List column: {' / '.join(names)}")


def product_key(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def translate_formula(value: str, src, dst) -> str:
    return Translator(value, origin=src.coordinate).translate_formula(dst.coordinate)


def clone_row_format_and_formulas(ws, src_row: int, dst_row: int) -> None:
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height
    for col in range(1, ws.max_column + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        dst._style = copy(src._style)
        if src.has_style:
            dst.font = copy(src.font)
            dst.fill = copy(src.fill)
            dst.border = copy(src.border)
            dst.alignment = copy(src.alignment)
            dst.number_format = src.number_format
            dst.protection = copy(src.protection)
        if src.data_type == "f" and src.value:
            dst.value = translate_formula(src.value, src, dst)


def copy_cell_value(src, dst) -> None:
    if src.data_type == "f" and src.value:
        dst.value = translate_formula(src.value, src, dst)
    else:
        dst.value = src.value


def last_non_empty_row(ws) -> int:
    for row in range(ws.max_row, 1, -1):
        if any(ws.cell(row, col).value is not None for col in range(1, ws.max_column + 1)):
            return row
    return 2


def update_lager2(kalkulation: Path, stock_list: Path, output: Path, overwrite: bool) -> dict[str, int]:
    if overwrite:
        output = kalkulation

    wb = openpyxl.load_workbook(kalkulation)
    stock_wb = openpyxl.load_workbook(stock_list, data_only=True)

    lager = wb["Lager"]
    target = wb["Lager (2)"]
    stock = stock_wb["Stock List"]

    lager_by_id = {}
    for row in range(2, lager.max_row + 1):
        key = product_key(lager.cell(row, 14).value)
        if key:
            lager_by_id[key] = {
                "row": row,
                "g": lager.cell(row, 7).value,
                "name_fill": copy(lager.cell(row, 3).fill),
            }

    header_values = next(stock.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = {
        normalize_header(value): idx
        for idx, value in enumerate(header_values)
        if value is not None
    }
    stock_rows = [
        values
        for values in stock.iter_rows(min_row=2, values_only=True)
        if product_key(value_from_row(values, headers, "Article ID", "Product ID", "Product UD"))
    ]
    final_row = len(stock_rows) + 1

    if target.max_row < final_row:
        template_row = last_non_empty_row(target)
        for row in range(target.max_row + 1, final_row + 1):
            clone_row_format_and_formulas(target, template_row, row)
    elif target.max_row > final_row:
        target.delete_rows(final_row + 1, target.max_row - final_row)

    matched_count = 0
    blue_count = 0

    for idx, values in enumerate(stock_rows, start=2):
        article_id = value_from_row(values, headers, "Article ID", "Product ID", "Product UD")
        name = value_from_row(values, headers, "Product name")
        size = value_from_row(values, headers, "Size")
        alcohol = value_from_row(values, headers, "Alcohol %")
        ean = value_from_row(values, headers, "EAN", "GTIN")
        bottles = value_from_row(values, headers, "Bottles per case")
        internal_price = value_from_row(values, headers, "Internal price", "Price/btl")
        sam_qty = value_from_row(values, headers, "SAM - Quantity Available", "SAM - Quantity in Hand", "Stock SAM")
        boz_qty = value_from_row(values, headers, "BOZ quantity available", "BOZ - Quantity Available", "Stock BOZ")
        key = product_key(article_id)
        matched = lager_by_id.get(key)

        target.cell(idx, 1).value = sam_qty
        target.cell(idx, 2).value = boz_qty
        target.cell(idx, 3).value = name
        target.cell(idx, 4).value = alcohol
        target.cell(idx, 5).value = bottles
        target.cell(idx, 6).value = size
        target.cell(idx, 7).value = matched["g"] if matched else None
        target.cell(idx, 14).value = article_id
        target.cell(idx, 15).value = internal_price
        target.cell(idx, 16).value = 0
        target.cell(idx, 18).value = ean
        if matched:
            source_row = matched["row"]
            copy_cell_value(lager.cell(source_row, 26), target.cell(idx, 26))
            copy_cell_value(lager.cell(source_row, 31), target.cell(idx, 31))
            copy_cell_value(lager.cell(source_row, 37), target.cell(idx, 37))
        else:
            target.cell(idx, 26).value = None
            target.cell(idx, 31).value = None
            target.cell(idx, 37).value = None

        if matched:
            matched_count += 1
            target.cell(idx, 3).fill = copy(matched["name_fill"])
        else:
            blue_count += 1
            target.cell(idx, 3).fill = copy(BLUE_FILL)

    wb.save(output)
    return {"updated_rows": len(stock_rows), "matched_rows": matched_count, "blue_rows": blue_count}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Update Kalkulation Lager (2) from Stock List.")
    parser.add_argument("--kalkulation", required=True, type=Path)
    parser.add_argument("--stock-list", required=True, type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--overwrite", action="store_true")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output = args.output or args.kalkulation.with_name("Kalkulation - Lager 2 updated.xlsx")
    stats = update_lager2(args.kalkulation, args.stock_list, output, args.overwrite)
    print(output)
    for key, value in stats.items():
        print(f"{key}: {value}")


if __name__ == "__main__":
    main()
