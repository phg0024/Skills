#!/usr/bin/env python3
import argparse
import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


BLUE_FILL = PatternFill(fill_type="solid", fgColor="FF0070C0")


TARGET_COLUMNS = {
    1: "stock",
    2: "name",
    3: "volume",
    4: "case_size",
    5: "size",
    6: "type",
    13: "product_id",
    14: "price_bottle",
    16: "currency",
    17: "ean",
}

HEADER_SYNONYMS = {
    "product_id": ["ProductID", "Product ID", "Item code", "Item Code", "Art-No Lieferant", "Supplier ID"],
    "stock": ["Stock", "Available", "avlbl cs", "Btls Available"],
    "name": ["Description", "ProductName", "Product Name", "Produktname / product name"],
    "volume": ["Vol %", "% Vol", "Vol", "Vol ", "Volume"],
    "case_size": ["Case size", "btl/Case", "Btl/case", "Qty/Case", "je Krt /\nper cs", "je Krt / per cs"],
    "size": ["size", "Size", "Inh / lt"],
    "type": ["Main Category", "Category", "Type", "Family"],
    "price_bottle": ["Price bottle", "Price/btl", "Supplier Price"],
    "unit_price": ["Unit price", "Price case", "Original Price"],
    "currency": ["Currency", "Währung"],
    "ean": ["EAN", "EAN code sales unit", "barcode", "Barcode"],
}


def norm(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def clean(value):
    if value == "":
        return None
    return value


def normalize_currency(value):
    text = norm(value)
    if text == "€":
        return "EUR"
    return clean(value)


def header_key(value):
    return norm(value).lower().replace("\n", " ")


def is_yellow(cell):
    fill = cell.fill
    if fill.fill_type is None:
        return False
    candidates = []
    for color in (fill.fgColor, fill.start_color):
        rgb = getattr(color, "rgb", None)
        indexed = getattr(color, "indexed", None)
        if rgb:
            candidates.append(str(rgb).upper())
        if indexed is not None:
            candidates.append(f"INDEXED:{indexed}")
    yellow_codes = {"FFFFFF00", "FFFF00", "FFFFEB9C", "FFFFFF99", "INDEXED:6", "INDEXED:13"}
    return any(candidate in yellow_codes for candidate in candidates)


def parse_volume(description):
    match = re.search(r"(\d+(?:[.,]\d+)?)\s*%", description or "")
    if not match:
        return None
    return float(match.group(1).replace(",", "."))


def parse_size_liters(description):
    matches = re.findall(r"(?i)(\d+(?:[.,]\d+)?)\s*(l|lt|ltr|cl|ml)\b", description or "")
    if not matches:
        return None
    value_text, unit = matches[-1]
    value = float(value_text.replace(",", "."))
    unit = unit.lower()
    if unit == "cl":
        return value / 100
    if unit == "ml":
        return value / 1000
    return value


def find_header(ws, requested_row=None):
    if requested_row:
        return requested_row
    for row in range(1, min(ws.max_row, 40) + 1):
        values = [header_key(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)]
        if any(value in {"productid", "product id", "item code"} for value in values):
            return row
    raise ValueError(f"Could not find a header row with ProductID/Item code in sheet {ws.title!r}")


def build_header_index(ws, header_row):
    index = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(header_row, col).value
        if value is not None:
            index[header_key(value)] = col
    return index


def find_col(index, field, explicit=None):
    if explicit:
        key = header_key(explicit)
        if key not in index:
            raise ValueError(f"Explicit header {explicit!r} was not found")
        return index[key]
    for synonym in HEADER_SYNONYMS[field]:
        key = header_key(synonym)
        if key in index:
            return index[key]
    return None


def source_sheet(workbook, requested=None):
    if requested:
        return workbook[requested]
    for ws in workbook.worksheets:
        try:
            find_header(ws)
            return ws
        except ValueError:
            continue
    raise ValueError("Could not find a source sheet with ProductID/Item code headers")


def load_preserved(target_ws):
    preserved = {
        "name": {},
        "type": {},
        "y": {},
        "ad": {},
        "b_fill": {},
        "yellow": set(),
        "empty_b_fill": None,
    }
    for row in range(2, target_ws.max_row + 1):
        product_id = norm(target_ws.cell(row, 13).value)
        if not product_id:
            continue
        preserved["name"][product_id] = target_ws.cell(row, 2).value
        preserved["type"][product_id] = target_ws.cell(row, 6).value
        preserved["y"][product_id] = target_ws.cell(row, 25).value
        preserved["ad"][product_id] = target_ws.cell(row, 30).value
        b_cell = target_ws.cell(row, 2)
        preserved["b_fill"][product_id] = copy(b_cell.fill)
        if is_yellow(b_cell):
            preserved["yellow"].add(product_id)
        elif preserved["empty_b_fill"] is None:
            preserved["empty_b_fill"] = copy(b_cell.fill)
    if preserved["empty_b_fill"] is None:
        preserved["empty_b_fill"] = PatternFill()
    return preserved


def name_fill(preserved, product_id):
    if product_id in preserved["b_fill"]:
        return copy(preserved["b_fill"][product_id])
    return copy(BLUE_FILL)


def read_products(ws, header_row, columns):
    products = []
    for row in range(header_row + 1, ws.max_row + 1):
        product_id = norm(ws.cell(row, columns["product_id"]).value)
        if not product_id:
            continue
        name = ws.cell(row, columns["name"]).value if columns["name"] else None
        case_size = ws.cell(row, columns["case_size"]).value if columns["case_size"] else None
        price_bottle = ws.cell(row, columns["price_bottle"]).value if columns["price_bottle"] else None
        unit_price = ws.cell(row, columns["unit_price"]).value if columns["unit_price"] else None
        if price_bottle is None and unit_price is not None and case_size:
            price_bottle = unit_price / case_size
        products.append(
            {
                "product_id": product_id,
                "stock": ws.cell(row, columns["stock"]).value if columns["stock"] else None,
                "name": name,
                "volume": ws.cell(row, columns["volume"]).value if columns["volume"] else parse_volume(name),
                "case_size": case_size,
                "size": ws.cell(row, columns["size"]).value if columns["size"] else parse_size_liters(name),
                "type": ws.cell(row, columns["type"]).value if columns["type"] else None,
                "price_bottle": price_bottle,
                "currency": normalize_currency(ws.cell(row, columns["currency"]).value) if columns["currency"] else None,
                "ean": clean(ws.cell(row, columns["ean"]).value) if columns["ean"] else None,
            }
        )
    return products


def update_tab(source_path, target_path, tab, output_path, source_sheet_name=None, header_row=None, explicit_headers=None):
    source_wb = load_workbook(source_path, data_only=True, read_only=True)
    source_ws = source_sheet(source_wb, source_sheet_name)
    header_row = find_header(source_ws, header_row)
    index = build_header_index(source_ws, header_row)
    columns = {
        field: find_col(index, field, explicit_headers.get(field))
        for field in HEADER_SYNONYMS
    }
    if not columns["product_id"]:
        raise ValueError("No source ProductID column found")

    target_wb = load_workbook(target_path, data_only=False, keep_links=False, rich_text=False)
    target_ws = target_wb[tab]
    preserved = load_preserved(target_ws)
    products = read_products(source_ws, header_row, columns)

    for target_row, product in enumerate(products, start=2):
        product_id = product["product_id"]
        product["name"] = preserved["name"].get(product_id) or product["name"]
        product["type"] = preserved["type"].get(product_id) or product["type"]
        for column, key in TARGET_COLUMNS.items():
            target_ws.cell(target_row, column).value = product[key]
        target_ws.cell(target_row, 25).value = preserved["y"].get(product_id)
        target_ws.cell(target_row, 30).value = preserved["ad"].get(product_id)
        target_ws.cell(target_row, 2).fill = name_fill(preserved, product_id)

    first_empty_row = len(products) + 2
    for row in range(first_empty_row, target_ws.max_row + 1):
        for column in TARGET_COLUMNS:
            target_ws.cell(row, column).value = None
        target_ws.cell(row, 25).value = None
        target_ws.cell(row, 30).value = None
        target_ws.cell(row, 2).fill = copy(preserved["empty_b_fill"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    target_wb.save(output_path)

    old_ids = set(preserved["b_fill"])
    source_ids = {product["product_id"] for product in products}
    new_ids = source_ids - old_ids
    return {
        "updated_rows": len(products),
        "preserved_y_ad": len(source_ids & set(preserved["y"])),
        "new_products": len(new_ids),
        "blue_products": len(new_ids),
        "yellow_preserved": len(source_ids & preserved["yellow"]),
        "removed_products": len(old_ids - source_ids),
        "output": str(output_path),
    }


def main():
    parser = argparse.ArgumentParser(description="Refresh one Kalkulation supplier price-list tab by ProductID.")
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--target", required=True, type=Path)
    parser.add_argument("--tab", required=True)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--source-sheet")
    parser.add_argument("--header-row", type=int)
    for field in HEADER_SYNONYMS:
        parser.add_argument(f"--{field.replace('_', '-')}-header")
    args = parser.parse_args()

    explicit_headers = {
        field: getattr(args, f"{field}_header")
        for field in HEADER_SYNONYMS
    }
    summary = update_tab(
        source_path=args.source,
        target_path=args.target,
        tab=args.tab,
        output_path=args.output,
        source_sheet_name=args.source_sheet,
        header_row=args.header_row,
        explicit_headers=explicit_headers,
    )
    for key, value in summary.items():
        print(f"{key}: {value}")


if __name__ == "__main__":
    main()
