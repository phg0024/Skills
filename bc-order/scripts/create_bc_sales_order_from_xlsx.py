#!/usr/bin/env python3
"""Validate, create, or safely resume a draft BC sales order from an XLSX."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any
from collections.abc import Iterable


DEFAULT_CONNECTOR = Path("/Users/ph/Documents/Silver Spirits Project Codex")
MONEY_TOLERANCE = Decimal("0.05")


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Validate, create, or resume a draft Business Central sales order "
            "from a Galaxus/Digitec XLSX order. Writes require --execute."
        )
    )
    parser.add_argument("xlsx_file", type=Path)
    parser.add_argument("--connector-dir", type=Path, default=DEFAULT_CONNECTOR)
    parser.add_argument("--customer-number", default="80324")
    parser.add_argument("--location-code", default="SAM")
    parser.add_argument("--external-doc-no")
    parser.add_argument("--sales-order-number")
    parser.add_argument("--order-date")
    parser.add_argument("--execute", action="store_true")
    args = parser.parse_args()

    sys.path.insert(0, str(args.connector_dir / "src"))
    from business_central import BusinessCentralClient, BusinessCentralConfig

    try:
        source_path = args.xlsx_file.expanduser().resolve()
        if not source_path.is_file():
            raise RuntimeError(f"Source workbook does not exist: {source_path}")
        if args.order_date:
            date.fromisoformat(args.order_date)

        rows = read_order_rows(source_path)
        external_doc_no = clean_text(args.external_doc_no) or infer_external_doc_no(source_path)
        currency = one_value(rows, "currency").upper()
        client = BusinessCentralClient(
            BusinessCentralConfig.from_env(args.connector_dir.expanduser().resolve() / ".env")
        )

        customer = resolve_single(
            client,
            "customers",
            f"number eq '{escape_odata(args.customer_number)}'",
            f"customer number {args.customer_number}",
        )
        location = resolve_single(
            client,
            "locations",
            f"code eq '{escape_odata(args.location_code)}'",
            f"location code {args.location_code}",
        )
        items = match_items(client, rows)

        existing_orders = find_orders(
            client,
            external_doc_no=external_doc_no,
            sales_order_number=args.sales_order_number,
        )
        if len(existing_orders) > 1:
            label = args.sales_order_number or external_doc_no
            raise RuntimeError(f"Order key {label!r} matched {len(existing_orders)} orders; refusing to choose.")

        order = existing_orders[0] if existing_orders else None
        if args.sales_order_number and order is None:
            raise RuntimeError(f"Sales order {args.sales_order_number} was not found.")
        if order:
            validate_order_header(
                order,
                customer_number=args.customer_number,
                external_doc_no=external_doc_no,
                currency=currency,
                allow_blank_external=bool(args.sales_order_number),
            )

        source_counter = Counter(
            source_fingerprint(row, location_id=location["id"]) for row in rows
        )
        existing_lines = fetch_lines(client, order["id"]) if order else []
        existing_counter = item_line_counter(existing_lines)
        unexpected = existing_counter - source_counter
        if unexpected:
            raise RuntimeError(
                "The target draft contains item lines not present in the source: "
                + format_counter(unexpected)
            )
        missing_counter = source_counter - existing_counter
        source_total = sum((row["total"] for row in rows), Decimal("0"))
        summary: dict[str, Any] = {
            "mode": "execute" if args.execute else "dry_run",
            "source": str(source_path),
            "externalDocumentNumber": external_doc_no,
            "customerNumber": customer.get("number"),
            "customerName": customer.get("displayName"),
            "locationCode": location.get("code"),
            "currencyCode": currency,
            "sourceLineCount": len(rows),
            "sourceTotalExcludingTax": decimal_json(source_total),
            "matchedUniqueItemCount": len(items),
            "existingOrderNumber": order.get("number") if order else None,
            "existingItemLineCount": sum(existing_counter.values()),
            "missingItemLineCount": sum(missing_counter.values()),
        }

        if not args.execute:
            print(json.dumps({"status": "validated_dry_run", **summary}, indent=2, ensure_ascii=False))
            return 0

        if order:
            order = ensure_external_document_number(client, order, external_doc_no)
        else:
            order = create_order_safely(
                client,
                customer_number=args.customer_number,
                external_doc_no=external_doc_no,
                currency=currency,
                order_date=args.order_date,
            )
            validate_order_header(
                order,
                customer_number=args.customer_number,
                external_doc_no=external_doc_no,
                currency=currency,
            )

        result = add_missing_lines_and_verify(
            client,
            order,
            rows,
            customer_number=args.customer_number,
            external_doc_no=external_doc_no,
            currency=currency,
            location_id=location["id"],
            source_total=source_total,
        )
        print(json.dumps({"status": "created_or_resumed", **summary, **result}, indent=2, ensure_ascii=False))
        return 0
    except Exception as exc:
        parser.exit(status=1, message=f"{exc}\n")


def read_order_rows(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required. Use the bundled Codex Python runtime.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        raw_headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not raw_headers:
            raise RuntimeError(f"Workbook has no header row: {path}")
        headers = [clean_text(value) for value in raw_headers]
        required = [
            "Supplier item number",
            "Item name",
            "Price per unit",
            "Quantity",
            "Total position",
            "Invoice currency",
        ]
        duplicate_headers = sorted(name for name, count in Counter(headers).items() if name and count > 1)
        if duplicate_headers:
            raise RuntimeError(f"Duplicate spreadsheet columns: {', '.join(duplicate_headers)}")
        positions = {header: index for index, header in enumerate(headers) if header}
        missing = [column for column in required if column not in positions]
        if missing:
            raise RuntimeError(f"Missing required spreadsheet columns: {', '.join(missing)}")

        rows: list[dict[str, Any]] = []
        for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            number = clean_identifier(cell_value(values, positions["Supplier item number"]))
            if not number:
                continue
            quantity = parse_decimal(cell_value(values, positions["Quantity"]), f"row {row_number} quantity")
            unit_price = parse_decimal(
                cell_value(values, positions["Price per unit"]), f"row {row_number} unit price"
            )
            total = parse_decimal(
                cell_value(values, positions["Total position"]), f"row {row_number} total"
            )
            currency = clean_text(cell_value(values, positions["Invoice currency"])).upper()
            if quantity <= 0:
                raise RuntimeError(f"Row {row_number}: quantity must be positive.")
            if unit_price < 0 or total < 0:
                raise RuntimeError(f"Row {row_number}: price and total cannot be negative.")
            calculated = (unit_price * quantity).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            if abs(calculated - total) > MONEY_TOLERANCE:
                raise RuntimeError(
                    f"Row {row_number}: quantity × unit price is {calculated}, but total is {total}."
                )
            if not currency:
                raise RuntimeError(f"Row {row_number}: invoice currency is blank.")
            rows.append(
                {
                    "row": row_number,
                    "number": number,
                    "description": clean_text(cell_value(values, positions["Item name"])),
                    "unitPrice": unit_price,
                    "quantity": quantity,
                    "total": total,
                    "currency": currency,
                }
            )
    finally:
        workbook.close()

    if not rows:
        raise RuntimeError(f"No order rows found in {path}")
    return rows


def cell_value(values: tuple[Any, ...], index: int) -> Any:
    return values[index] if index < len(values) else None


def parse_decimal(value: Any, label: str) -> Decimal:
    if isinstance(value, Decimal):
        result = value
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        result = Decimal(str(value))
    else:
        text = clean_text(value).replace("\u00a0", "").replace(" ", "").replace("'", "")
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        elif "," in text:
            text = text.replace(",", ".")
        try:
            result = Decimal(text)
        except (InvalidOperation, ValueError) as exc:
            raise RuntimeError(f"Invalid {label}: {value!r}") from exc
    if not result.is_finite():
        raise RuntimeError(f"Invalid {label}: {value!r}")
    return result


def clean_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def clean_identifier(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(value)


def infer_external_doc_no(path: Path) -> str:
    matches = list(dict.fromkeys(re.findall(r"\d{6,}", path.stem)))
    if not matches:
        raise RuntimeError("Could not infer external document number; pass --external-doc-no.")
    if len(matches) > 1:
        raise RuntimeError(
            f"Filename contains multiple possible document numbers {matches}; pass --external-doc-no."
        )
    return matches[0]


def one_value(rows: list[dict[str, Any]], key: str) -> str:
    values = {clean_text(row.get(key)) for row in rows if clean_text(row.get(key))}
    if len(values) != 1:
        raise RuntimeError(f"Expected exactly one {key}, found {sorted(values)}")
    return next(iter(values))


def resolve_single(client: Any, resource: str, filter_expr: str, label: str) -> dict[str, Any]:
    select_fields = {
        "customers": "id,number,displayName",
        "locations": "id,code,displayName",
    }.get(resource, "id,number,code,displayName")
    rows = client.get_all(
        client.company_path(resource),
        query={"$filter": filter_expr, "$select": select_fields},
    )
    if len(rows) != 1:
        raise RuntimeError(f"Expected exactly one {label}, found {len(rows)}")
    return rows[0]


def chunks(values: list[str], size: int) -> Iterable[list[str]]:
    for start in range(0, len(values), size):
        yield values[start : start + size]


def match_items(client: Any, rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    numbers = sorted({row["number"] for row in rows})
    matches: dict[str, list[dict[str, Any]]] = {number: [] for number in numbers}
    for group in chunks(numbers, 20):
        filter_expr = " or ".join(f"number eq '{escape_odata(number)}'" for number in group)
        found = client.get_all(
            client.company_path("items"),
            query={"$filter": filter_expr, "$select": "id,number,displayName"},
        )
        for item in found:
            number = clean_text(item.get("number"))
            if number in matches:
                matches[number].append(item)
    problems = [f"{number} ({len(found)} matches)" for number, found in matches.items() if len(found) != 1]
    if problems:
        raise RuntimeError("Every source item must match exactly one BC item: " + ", ".join(problems))
    return {number: found[0] for number, found in matches.items()}


def find_orders(
    client: Any,
    *,
    external_doc_no: str,
    sales_order_number: str | None,
) -> list[dict[str, Any]]:
    if sales_order_number:
        filter_expr = f"number eq '{escape_odata(sales_order_number)}'"
    else:
        filter_expr = f"externalDocumentNumber eq '{escape_odata(external_doc_no)}'"
    return client.get_all(client.company_path("salesOrders"), query={"$filter": filter_expr})


def validate_order_header(
    order: dict[str, Any],
    *,
    customer_number: str,
    external_doc_no: str,
    currency: str,
    allow_blank_external: bool = False,
) -> None:
    status = clean_text(order.get("status")).casefold()
    if status != "draft":
        raise RuntimeError(
            f"Sales order {order.get('number') or order.get('id')} is {order.get('status')!r}, not Draft."
        )
    actual_customer = clean_text(order.get("customerNumber"))
    if actual_customer != clean_text(customer_number):
        raise RuntimeError(f"Target order customer is {actual_customer!r}, expected {customer_number!r}.")
    actual_currency = clean_text(order.get("currencyCode")).upper()
    if actual_currency != currency.upper():
        raise RuntimeError(f"Target order currency is {actual_currency!r}, expected {currency!r}.")
    actual_external = clean_text(order.get("externalDocumentNumber"))
    if actual_external != external_doc_no and not (allow_blank_external and not actual_external):
        raise RuntimeError(
            f"Target order external document is {actual_external!r}, expected {external_doc_no!r}."
        )


def ensure_external_document_number(client: Any, order: dict[str, Any], external_doc_no: str) -> dict[str, Any]:
    current = clean_text(order.get("externalDocumentNumber"))
    if current == external_doc_no:
        return order
    if current:
        raise RuntimeError(f"Target order already has external document number {current!r}.")
    path = client.company_path(f"salesOrders({order['id']})")
    try:
        updated = client.patch(
            path,
            {"externalDocumentNumber": external_doc_no},
            etag=order.get("@odata.etag") or "*",
        )
    except Exception:
        reread = client.get(path)
        if clean_text(reread.get("externalDocumentNumber")) != external_doc_no:
            raise
        return reread
    return updated or client.get(path)


def create_order_safely(
    client: Any,
    *,
    customer_number: str,
    external_doc_no: str,
    currency: str,
    order_date: str | None,
) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "customerNumber": customer_number,
        "externalDocumentNumber": external_doc_no,
        "currencyCode": currency,
    }
    if order_date:
        payload["orderDate"] = order_date
        payload["postingDate"] = order_date
    try:
        created = client.post(client.company_path("salesOrders"), payload)
    except Exception:
        recovered = find_orders(client, external_doc_no=external_doc_no, sales_order_number=None)
        if len(recovered) != 1:
            raise
        created = recovered[0]
    if not created.get("id"):
        raise RuntimeError("Business Central did not return an order id after header creation.")
    return created


def fetch_lines(client: Any, order_id: str) -> list[dict[str, Any]]:
    return client.get_all(client.company_path(f"salesOrders({order_id})/salesOrderLines"))


def normalized_decimal(value: Any, places: str = "0.0001") -> str:
    parsed = parse_decimal(value or 0, "Business Central numeric value")
    return format(parsed.quantize(Decimal(places), rounding=ROUND_HALF_UP), "f")


def source_fingerprint(row: dict[str, Any], *, location_id: str) -> tuple[str, str, str, str]:
    return (
        clean_text(row["number"]),
        normalized_decimal(row["quantity"]),
        normalized_decimal(row["unitPrice"]),
        clean_text(location_id),
    )


def bc_line_fingerprint(line: dict[str, Any]) -> tuple[str, str, str, str]:
    return (
        clean_text(line.get("lineObjectNumber")),
        normalized_decimal(line.get("quantity")),
        normalized_decimal(line.get("unitPrice")),
        clean_text(line.get("locationId")),
    )


def item_line_counter(lines: list[dict[str, Any]]) -> Counter[tuple[str, str, str, str]]:
    return Counter(
        bc_line_fingerprint(line)
        for line in lines
        if clean_text(line.get("lineType")).casefold() == "item"
    )


def format_counter(counter: Counter[tuple[str, str, str, str]]) -> str:
    return "; ".join(
        f"{number} qty={quantity} price={price} location={location or '<blank>'} ×{count}"
        for (number, quantity, price, location), count in sorted(counter.items())
    )


def add_missing_lines_and_verify(
    client: Any,
    order: dict[str, Any],
    rows: list[dict[str, Any]],
    *,
    customer_number: str,
    external_doc_no: str,
    currency: str,
    location_id: str,
    source_total: Decimal,
) -> dict[str, Any]:
    order_id = order["id"]
    source_counter = Counter(source_fingerprint(row, location_id=location_id) for row in rows)
    existing_lines = fetch_lines(client, order_id)
    existing_counter = item_line_counter(existing_lines)
    unexpected = existing_counter - source_counter
    if unexpected:
        raise RuntimeError("Target draft has unexpected item lines: " + format_counter(unexpected))

    created_count = 0
    line_path = client.company_path(f"salesOrders({order_id})/salesOrderLines")
    for row in rows:
        fingerprint = source_fingerprint(row, location_id=location_id)
        if existing_counter[fingerprint] >= source_counter[fingerprint]:
            continue
        before_count = existing_counter[fingerprint]
        payload = {
            "lineType": "Item",
            "lineObjectNumber": row["number"],
            "quantity": float(row["quantity"]),
            "unitPrice": float(row["unitPrice"]),
            "locationId": location_id,
        }
        try:
            client.post(line_path, payload)
        except Exception:
            after_failure = item_line_counter(fetch_lines(client, order_id))
            if after_failure[fingerprint] != before_count + 1:
                raise
            existing_counter = after_failure
        else:
            existing_counter[fingerprint] += 1
        created_count += 1

    verified_order = client.get(client.company_path(f"salesOrders({order_id})"))
    validate_order_header(
        verified_order,
        customer_number=customer_number,
        external_doc_no=external_doc_no,
        currency=currency,
    )
    verified_lines = fetch_lines(client, order_id)
    verified_counter = item_line_counter(verified_lines)
    if verified_counter != source_counter:
        missing = source_counter - verified_counter
        extra = verified_counter - source_counter
        raise RuntimeError(
            "Saved item lines do not exactly match the source. "
            f"Missing: {format_counter(missing) or 'none'}. Extra: {format_counter(extra) or 'none'}."
        )

    bc_total_raw = verified_order.get("totalAmountExcludingTax")
    if bc_total_raw not in (None, ""):
        bc_total = parse_decimal(bc_total_raw, "Business Central total excluding tax")
        if abs(bc_total - source_total) > MONEY_TOLERANCE:
            raise RuntimeError(
                f"Business Central total excluding tax is {bc_total}, source total is {source_total}."
            )
    else:
        bc_total = None

    non_item_lines = [
        line for line in verified_lines if clean_text(line.get("lineType")).casefold() != "item"
    ]
    if non_item_lines:
        raise RuntimeError(f"Saved order contains {len(non_item_lines)} unexpected non-item line(s).")
    return {
        "orderId": order_id,
        "orderNumber": verified_order.get("number"),
        "bcStatus": verified_order.get("status"),
        "createdLineCount": created_count,
        "verifiedItemLineCount": sum(verified_counter.values()),
        "bcTotalExcludingTax": decimal_json(bc_total) if bc_total is not None else None,
    }


def decimal_json(value: Decimal) -> float:
    return float(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def escape_odata(value: Any) -> str:
    return clean_text(value).replace("'", "''")


if __name__ == "__main__":
    raise SystemExit(main())
