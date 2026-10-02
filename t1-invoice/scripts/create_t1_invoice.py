#!/usr/bin/env python3
"""Validate and create/resume a draft CHF T1 sales invoice for customer 80211."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader


CUSTOMER_NUMBER = "80211"
LOCATION_CODE = "BOZ"
CURRENCY_CODE = "CHF"
TRANSPORT_ITEM_NUMBER = "599999"
CURRENCY_QUANTUM = Decimal("0.01")
MONEY_TOLERANCE = Decimal("0.05")
PRICE_TOLERANCE = Decimal("0.0001")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Validate or create the unposted CHF T1 invoice for customer 80211."
    )
    parser.add_argument("source_xlsx", type=Path)
    parser.add_argument("emcs_pdf", type=Path)
    parser.add_argument(
        "--exclude-item",
        action="append",
        required=True,
        help="Source/BC item number for each product appearing on the EMCS PDF.",
    )
    parser.add_argument("--invoice-date", help="Invoice/posting date YYYY-MM-DD; defaults to today.")
    parser.add_argument(
        "--external-document-number",
        help="Override the EMCS invoice number used as the duplicate-safe external reference.",
    )
    parser.add_argument(
        "--emcs-product-count",
        type=int,
        help="Override the number of EMCS body product records if text extraction cannot detect it.",
    )
    parser.add_argument(
        "--connector-dir",
        type=Path,
        default=None,
        help="Silver Spirits workspace containing src/business_central and .env.",
    )
    parser.add_argument("--execute", action="store_true", help="Create or repair the Draft invoice.")
    args = parser.parse_args()

    try:
        source_rows = read_source_rows(args.source_xlsx.expanduser().resolve())
        pdf_text = read_pdf_text(args.emcs_pdf.expanduser().resolve())
        excluded_numbers = clean_exclusion_numbers(args.exclude_item)
        verify_emcs_exclusions(source_rows, pdf_text, excluded_numbers, args.emcs_product_count)
        included_rows = [
            row
            for row in source_rows
            if row["number"] not in excluded_numbers
            and row["number"] != TRANSPORT_ITEM_NUMBER
        ]
        if not included_rows:
            raise RuntimeError("No invoice lines remain after EMCS and transport exclusions.")

        connector_dir = resolve_connector_dir(args.connector_dir)
        sys.path.insert(0, str(connector_dir / "src"))
        from business_central import BusinessCentralClient, BusinessCentralConfig

        client = BusinessCentralClient(BusinessCentralConfig.from_env(connector_dir / ".env"))
        customer = resolve_single(
            client,
            "customers",
            f"number eq '{escape_odata(CUSTOMER_NUMBER)}'",
            f"customer number {CUSTOMER_NUMBER}",
            select="id,number,displayName,country,currencyCode,blocked",
        )
        validate_customer(customer)
        location = resolve_single(
            client,
            "locations",
            f"code eq '{escape_odata(LOCATION_CODE)}'",
            f"location code {LOCATION_CODE}",
            select="id,code,displayName",
        )
        items = match_items(client, source_rows)
        last_costs = match_last_direct_costs(client, included_rows)
        prepared_rows = prepare_rows(included_rows, items, last_costs)

        invoice_date = args.invoice_date or date.today().isoformat()
        date.fromisoformat(invoice_date)
        external_number = args.external_document_number or infer_external_document_number(pdf_text)
        existing = find_invoice_by_external_number(client, external_number)
        if len(existing) > 1:
            raise RuntimeError(
                f"External document number {external_number!r} matched {len(existing)} invoices."
            )
        if existing:
            validate_invoice_header(existing[0], external_number)

        source_total = sum((row["sourceTotal"] for row in source_rows), Decimal("0"))
        expected_total = sum((row["expectedAmount"] for row in prepared_rows), Decimal("0"))
        summary: dict[str, Any] = {
            "mode": "execute" if args.execute else "dry_run",
            "sourceFile": str(args.source_xlsx.expanduser().resolve()),
            "emcsFile": str(args.emcs_pdf.expanduser().resolve()),
            "customerNumber": customer["number"],
            "customerName": customer["displayName"],
            "locationCode": location["code"],
            "currencyCode": CURRENCY_CODE,
            "invoiceDate": invoice_date,
            "externalDocumentNumber": external_number,
            "emcsProductCount": args.emcs_product_count or detect_emcs_product_count(pdf_text),
            "sourceLineCount": len(source_rows),
            "includedLineCount": len(prepared_rows),
            "excludedEmcsItemNumbers": sorted(excluded_numbers),
            "transportExcludedItemNumber": TRANSPORT_ITEM_NUMBER,
            "sourceTotalFromWorkbook": decimal_json(source_total),
            "expectedDraftTotalExcludingTax": decimal_json(expected_total),
            "expectedQuantity": decimal_json(sum((row["quantity"] for row in prepared_rows), Decimal("0"))),
            "existingInvoiceNumber": existing[0].get("number") if existing else None,
            "lines": [
                {
                    "itemNumber": row["number"],
                    "description": row["description"],
                    "quantity": decimal_json(row["quantity"]),
                    "lastDirectCostCHF": decimal_json(row["unitPrice"]),
                    "expectedAmountCHF": decimal_json(row["expectedAmount"]),
                    "lastCostDate": row["lastCostDate"],
                }
                for row in prepared_rows
            ],
        }
        if not args.execute:
            print(json.dumps({"status": "validated_dry_run", **summary}, indent=2, ensure_ascii=False))
            return 0

        invoice, resumed = find_or_create_invoice(
            client,
            external_number=external_number,
            invoice_date=invoice_date,
        )
        result = add_lines_and_verify(
            client,
            invoice,
            prepared_rows,
            location_id=location["id"],
            expected_total=expected_total,
            external_number=external_number,
        )
        print(
            json.dumps(
                {
                    "status": "created_or_resumed_draft",
                    **summary,
                    "resumedExistingDraft": resumed,
                    **result,
                },
                indent=2,
                ensure_ascii=False,
            )
        )
        return 0
    except Exception as exc:
        parser.exit(status=1, message=f"{exc}\n")


def read_source_rows(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        raise RuntimeError(f"Source workbook does not exist: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Lines" not in workbook.sheetnames:
            raise RuntimeError("Source workbook is missing the Lines sheet.")
        worksheet = workbook["Lines"]
        raw_headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        headers = [clean_text(value) for value in raw_headers or ()]
        required = ["No.", "Description", "Quantity"]
        missing = [name for name in required if name not in headers]
        if missing:
            raise RuntimeError("Missing source columns: " + ", ".join(missing))
        duplicate_headers = sorted(
            name for name, count in Counter(headers).items() if name and count > 1
        )
        if duplicate_headers:
            raise RuntimeError("Duplicate source columns: " + ", ".join(duplicate_headers))
        positions = {name: index for index, name in enumerate(headers) if name}
        has_source_cost = (
            "Direct Unit Cost Excl. VAT" in positions
            and "Line Amount Excl. VAT" in positions
        )
        rows: list[dict[str, Any]] = []
        seen: set[str] = set()
        for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(clean_text(value) for value in values[:3]):
                continue
            number = clean_identifier(cell_value(values, positions["No."]))
            description = clean_text(cell_value(values, positions["Description"]))
            if not number:
                raise RuntimeError(f"Source row {row_number} has no item number.")
            if number in seen:
                raise RuntimeError(f"Duplicate source item number {number} on row {row_number}.")
            seen.add(number)
            quantity = parse_decimal(
                cell_value(values, positions["Quantity"]),
                f"source row {row_number} quantity",
            )
            if quantity <= 0 or quantity != quantity.to_integral_value():
                raise RuntimeError(f"Source row {row_number}: quantity must be a positive whole number.")
            source_unit_cost = None
            source_total = None
            if has_source_cost:
                raw_unit = cell_value(values, positions["Direct Unit Cost Excl. VAT"])
                raw_total = cell_value(values, positions["Line Amount Excl. VAT"])
                if clean_text(raw_unit) or clean_text(raw_total):
                    source_unit_cost = parse_decimal(raw_unit, f"source row {row_number} direct unit cost")
                    source_total = parse_decimal(raw_total, f"source row {row_number} line amount")
                    if source_unit_cost < 0 or source_total < 0:
                        raise RuntimeError(f"Source row {row_number}: source costs cannot be negative.")
                    calculated = (quantity * source_unit_cost).quantize(
                        Decimal("0.01"), rounding=ROUND_HALF_UP
                    )
                    if abs(calculated - source_total) > MONEY_TOLERANCE:
                        raise RuntimeError(
                            f"Source row {row_number}: quantity × source cost is {calculated}, "
                            f"but line amount is {source_total}."
                        )
            rows.append(
                {
                    "sourceRow": row_number,
                    "number": number,
                    "description": description,
                    "quantity": quantity,
                    "sourceUnitCost": source_unit_cost,
                    "sourceTotal": source_total or Decimal("0"),
                }
            )
    finally:
        workbook.close()
    if not rows:
        raise RuntimeError("No source lines found.")
    return rows


def read_pdf_text(path: Path) -> str:
    if not path.is_file():
        raise RuntimeError(f"EMCS PDF does not exist: {path}")
    reader = PdfReader(path)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if not text.strip():
        raise RuntimeError("EMCS PDF contained no extractable text.")
    return text


def clean_exclusion_numbers(values: list[str]) -> set[str]:
    cleaned = {clean_identifier(value) for value in values}
    if len(cleaned) != len(values):
        raise RuntimeError("Duplicate --exclude-item values were provided.")
    if TRANSPORT_ITEM_NUMBER in cleaned:
        raise RuntimeError("Do not pass transport item 599999 as an EMCS product exclusion; it is excluded automatically.")
    return cleaned


def verify_emcs_exclusions(
    source_rows: list[dict[str, Any]],
    pdf_text: str,
    excluded_numbers: set[str],
    product_count_override: int | None,
) -> None:
    source_by_number = {row["number"]: row for row in source_rows}
    missing = sorted(excluded_numbers - set(source_by_number))
    if missing:
        raise RuntimeError("EMCS exclusions not present in source workbook: " + ", ".join(missing))
    product_count = product_count_override or detect_emcs_product_count(pdf_text)
    if product_count <= 0:
        raise RuntimeError("Could not detect EMCS body product records; provide --emcs-product-count.")
    if len(excluded_numbers) != product_count:
        raise RuntimeError(
            f"Mapped {len(excluded_numbers)} EMCS products, but the PDF contains {product_count} body product records."
        )
    pdf_words = set(re.findall(r"[a-z0-9]+", pdf_text.casefold()))
    for number in sorted(excluded_numbers):
        key = first_distinctive_product_word(source_by_number[number]["description"])
        if key not in pdf_words:
            raise RuntimeError(
                f"Excluded source item {number} ({source_by_number[number]['description']!r}) "
                f"could not be verified in the EMCS PDF."
            )


def detect_emcs_product_count(pdf_text: str) -> int:
    matches = re.findall(r"(?m)^\s*A\s+(\d+)\s+(?=[A-Z0-9])", pdf_text)
    return len(set(matches))


def first_distinctive_product_word(description: str) -> str:
    noise = {
        "the", "yo", "years", "year", "lt", "liter", "liters", "bottle", "bottles",
        "carton", "cask", "gb", "uk", "us", "usa", "de", "it", "fr", "es", "mx",
        "pr", "jm", "bq", "united", "kingdom", "states", "america", "germany", "france",
        "spain", "italy", "mexico", "jamaica", "caribbean", "dominican", "republic",
        "puerto", "rico", "gin", "rum", "brandy", "whisky", "whiskey", "premium", "barrel",
    }
    for token in re.findall(r"[a-z0-9]+", description.casefold()):
        if len(token) >= 4 and token not in noise and not token.isdigit():
            return token
    raise RuntimeError(f"Could not identify a distinctive product word in {description!r}.")


def resolve_connector_dir(value: Path | None) -> Path:
    if value:
        connector_dir = value.expanduser().resolve()
    else:
        connector_dir = None
        for candidate in Path(__file__).resolve().parents:
            if (candidate / "src" / "business_central").is_dir() and (candidate / ".env").is_file():
                connector_dir = candidate
                break
        if connector_dir is None:
            connector_dir = Path.cwd().resolve()
    if not (connector_dir / "src" / "business_central").is_dir():
        raise RuntimeError(f"Business Central connector not found under {connector_dir}.")
    return connector_dir


def resolve_single(
    client: Any,
    resource: str,
    filter_expr: str,
    label: str,
    *,
    select: str,
) -> dict[str, Any]:
    rows = client.get_all(
        client.company_path(resource),
        query={"$filter": filter_expr, "$select": select},
    )
    if len(rows) != 1:
        raise RuntimeError(f"Expected exactly one {label}, found {len(rows)}.")
    return rows[0]


def validate_customer(customer: dict[str, Any]) -> None:
    blocked = clean_text(customer.get("blocked")).casefold()
    if blocked not in {"", " ", "_x0020_"}:
        raise RuntimeError(f"Customer {CUSTOMER_NUMBER} is blocked ({customer.get('blocked')!r}).")
    if clean_text(customer.get("currencyCode")).upper() != CURRENCY_CODE:
        raise RuntimeError(f"Customer {CUSTOMER_NUMBER} is not configured for CHF.")


def match_items(client: Any, rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    numbers = sorted({row["number"] for row in rows if row["number"] != TRANSPORT_ITEM_NUMBER})
    if TRANSPORT_ITEM_NUMBER in {row["number"] for row in rows}:
        numbers.append(TRANSPORT_ITEM_NUMBER)
    filter_expr = " or ".join(f"number eq '{escape_odata(number)}'" for number in numbers)
    found = client.get_all(
        client.company_path("items"),
        query={"$filter": filter_expr, "$select": "id,number,displayName,type,blocked"},
    )
    matches: dict[str, list[dict[str, Any]]] = {number: [] for number in numbers}
    for item in found:
        number = clean_identifier(item.get("number"))
        if number in matches:
            matches[number].append(item)
    problems = [
        f"{number} ({len(found_items)} matches)"
        for number, found_items in matches.items()
        if len(found_items) != 1
    ]
    if problems:
        raise RuntimeError("Every source item must match exactly one BC item: " + ", ".join(problems))
    return {number: found_items[0] for number, found_items in matches.items()}


def match_last_direct_costs(client: Any, rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    numbers = sorted({row["number"] for row in rows})
    filter_expr = " or ".join(f"No eq '{escape_odata(number)}'" for number in numbers)
    found = client.get_odata_all(
        "Artikelkarte_Excel",
        query={"$filter": filter_expr, "$select": "No,Last_Direct_Cost,Last_Date_Modified"},
    )
    matches: dict[str, list[dict[str, Any]]] = {number: [] for number in numbers}
    for item in found:
        number = clean_identifier(item.get("No"))
        if number in matches:
            matches[number].append(item)
    problems = [
        f"{number} ({len(found_items)} matches)"
        for number, found_items in matches.items()
        if len(found_items) != 1
    ]
    if problems:
        raise RuntimeError("Every invoice item must have one Last Direct Cost: " + ", ".join(problems))
    result: dict[str, dict[str, Any]] = {}
    for number, found_items in matches.items():
        raw_cost = parse_decimal(found_items[0].get("Last_Direct_Cost"), f"{number} Last Direct Cost")
        if raw_cost < 0:
            raise RuntimeError(f"Last Direct Cost for {number} is negative: {raw_cost}.")
        rounded_cost = raw_cost.quantize(CURRENCY_QUANTUM, rounding=ROUND_HALF_UP)
        result[number] = {**found_items[0], "lastDirectCost": rounded_cost}
    return result


def prepare_rows(
    rows: list[dict[str, Any]],
    items: dict[str, dict[str, Any]],
    costs: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    prepared: list[dict[str, Any]] = []
    for row in rows:
        cost = costs[row["number"]]
        unit_price = cost["lastDirectCost"]
        prepared.append(
            {
                **row,
                "bcItemId": items[row["number"]]["id"],
                "itemType": clean_text(items[row["number"]].get("type")),
                "unitPrice": unit_price,
                "expectedAmount": (row["quantity"] * unit_price).quantize(
                    Decimal("0.01"), rounding=ROUND_HALF_UP
                ),
                "lastCostDate": clean_text(cost.get("Last_Date_Modified")),
            }
        )
    return prepared


def infer_external_document_number(pdf_text: str) -> str:
    lines = pdf_text.splitlines()
    for index, line in enumerate(lines):
        match = re.search(r"invoice\s+number", line, re.IGNORECASE)
        if not match:
            continue
        candidates: list[str] = []
        same_line_tail = line[match.end():]
        candidates.extend(re.findall(r"[A-Z0-9][A-Z0-9-]{5,}", same_line_tail, re.IGNORECASE))
        for next_line in lines[index + 1 : index + 3]:
            candidates.extend(re.findall(r"[A-Z0-9][A-Z0-9-]{5,}", next_line, re.IGNORECASE))
            if candidates:
                # The EMCS layout can place the consignee excise number and invoice
                # number together. The invoice value is the last value in that row.
                return candidates[-1]
        if candidates:
            return candidates[-1]
    raise RuntimeError("No EMCS invoice number found; pass --external-document-number.")


def find_invoice_by_external_number(client: Any, external_number: str) -> list[dict[str, Any]]:
    return client.get_all(
        client.company_path("salesInvoices"),
        query={"$filter": f"externalDocumentNumber eq '{escape_odata(external_number)}'"},
    )


def find_or_create_invoice(
    client: Any,
    *,
    external_number: str,
    invoice_date: str,
) -> tuple[dict[str, Any], bool]:
    existing = find_invoice_by_external_number(client, external_number)
    if len(existing) > 1:
        raise RuntimeError(f"External document number {external_number!r} is ambiguous.")
    if existing:
        validate_invoice_header(existing[0], external_number)
        return existing[0], True
    payload = {
        "customerNumber": CUSTOMER_NUMBER,
        "invoiceDate": invoice_date,
        "postingDate": invoice_date,
        "currencyCode": CURRENCY_CODE,
        "externalDocumentNumber": external_number,
    }
    try:
        invoice = client.post(client.company_path("salesInvoices"), payload)
    except Exception as exc:
        recovered = find_invoice_by_external_number(client, external_number)
        if len(recovered) == 1:
            invoice = recovered[0]
        else:
            raise RuntimeError(f"Could not safely create the draft invoice header: {exc}") from exc
    validate_invoice_header(invoice, external_number)
    return invoice, False


def validate_invoice_header(invoice: dict[str, Any], external_number: str) -> None:
    if clean_text(invoice.get("status")).casefold() != "draft":
        raise RuntimeError(f"Invoice {invoice.get('number')} is not Draft.")
    if clean_identifier(invoice.get("customerNumber")) != CUSTOMER_NUMBER:
        raise RuntimeError(f"Invoice {invoice.get('number')} is not for customer {CUSTOMER_NUMBER}.")
    if clean_text(invoice.get("currencyCode")).upper() != CURRENCY_CODE:
        raise RuntimeError(f"Invoice {invoice.get('number')} is not in CHF.")
    if clean_text(invoice.get("externalDocumentNumber")) != external_number:
        raise RuntimeError("Invoice external document number does not match the source reference.")


def fetch_invoice_lines(client: Any, invoice_id: str) -> list[dict[str, Any]]:
    return client.get_all(
        client.company_path(f"salesInvoices({invoice_id})/salesInvoiceLines"),
        query={"$top": 1000},
    )


def add_lines_and_verify(
    client: Any,
    invoice: dict[str, Any],
    rows: list[dict[str, Any]],
    *,
    location_id: str,
    expected_total: Decimal,
    external_number: str,
) -> dict[str, Any]:
    invoice_id = invoice["id"]
    line_path = client.company_path(f"salesInvoices({invoice_id})/salesInvoiceLines")
    existing_lines = fetch_invoice_lines(client, invoice_id)
    existing_by_number = index_item_lines(existing_lines)
    expected_numbers = {row["number"] for row in rows}
    unexpected = sorted(set(existing_by_number) - expected_numbers - {TRANSPORT_ITEM_NUMBER})
    if unexpected:
        raise RuntimeError("Draft contains unexpected item lines: " + ", ".join(unexpected))

    removed_transport = False
    transport_lines = [
        line for line in existing_lines
        if clean_identifier(line.get("lineObjectNumber")) == TRANSPORT_ITEM_NUMBER
    ]
    if len(transport_lines) > 1:
        raise RuntimeError("Draft contains duplicate transport lines.")
    if transport_lines:
        delete_line_safely(client, line_path, transport_lines[0])
        removed_transport = True

    for row in rows:
        current = existing_by_number.get(row["number"])
        if current is None:
            add_line_safely(client, line_path, row, location_id)
        else:
            validate_or_repair_existing_line(client, line_path, current, row, location_id)

    if removed_transport:
        refresh_invoice_totals(client, invoice_id, invoice.get("invoiceDate"))

    verified_invoice = client.get(client.company_path(f"salesInvoices({invoice_id})"))
    validate_invoice_header(verified_invoice, external_number)
    verified_lines = fetch_invoice_lines(client, invoice_id)
    verified_by_number = index_item_lines(verified_lines)
    if set(verified_by_number) != expected_numbers:
        raise RuntimeError(
            "Draft line set mismatch after update. "
            f"Missing={sorted(expected_numbers - set(verified_by_number))}, "
            f"unexpected={sorted(set(verified_by_number) - expected_numbers)}."
        )

    line_total = Decimal("0")
    for row in rows:
        line = verified_by_number[row["number"]]
        validate_line_values(line, row, location_id)
        line_total += parse_decimal(line.get("amountExcludingTax"), f"line {row['number']} amount")
    line_total = line_total.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    header_total = parse_decimal(
        verified_invoice.get("totalAmountExcludingTax"), "invoice total"
    ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if abs(line_total - expected_total) > MONEY_TOLERANCE:
        raise RuntimeError(f"Verified line total is {line_total}, expected {expected_total}.")
    if abs(header_total - expected_total) > MONEY_TOLERANCE:
        raise RuntimeError(f"BC header total is {header_total}, expected {expected_total}.")
    return {
        "invoiceNumber": verified_invoice.get("number"),
        "invoiceId": verified_invoice.get("id"),
        "statusAfter": verified_invoice.get("status"),
        "lineCount": len(verified_by_number),
        "totalQuantity": decimal_json(sum((row["quantity"] for row in rows), Decimal("0"))),
        "totalAmountExcludingTax": decimal_json(header_total),
        "totalAmountIncludingTax": verified_invoice.get("totalAmountIncludingTax"),
        "removedTransport": removed_transport,
        "posted": False,
    }


def index_item_lines(lines: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for line in lines:
        if clean_text(line.get("lineType")).casefold() != "item":
            raise RuntimeError(f"Draft contains a non-item line ({line.get('lineType')!r}).")
        number = clean_identifier(line.get("lineObjectNumber"))
        if not number:
            raise RuntimeError("Draft contains an item line without an item number.")
        if number in result:
            raise RuntimeError(f"Draft contains duplicate item lines for {number}.")
        result[number] = line
    return result


def add_line_safely(client: Any, line_path: str, row: dict[str, Any], location_id: str) -> None:
    payload = {
        "lineType": "Item",
        "lineObjectNumber": row["number"],
        "quantity": float(row["quantity"]),
        "unitPrice": float(row["unitPrice"]),
        "locationId": location_id,
    }
    try:
        client.post(line_path, payload)
    except Exception as exc:
        refreshed = client.get_all(line_path, query={"$top": 1000})
        matches = [
            line for line in refreshed
            if clean_identifier(line.get("lineObjectNumber")) == row["number"]
            and parse_decimal(line.get("quantity"), f"line {row['number']} quantity") == row["quantity"]
        ]
        if len(matches) == 1:
            validate_line_values(matches[0], row, location_id)
            return
        raise RuntimeError(f"Could not safely create invoice line {row['number']}: {exc}") from exc


def delete_line_safely(client: Any, line_path: str, line: dict[str, Any]) -> None:
    try:
        client.delete(f"{line_path}({line['id']})")
    except Exception as exc:
        refreshed = client.get_all(line_path, query={"$top": 1000})
        if any(str(candidate.get("id")) == str(line.get("id")) for candidate in refreshed):
            raise RuntimeError(f"Could not safely remove line {line.get('lineObjectNumber')}: {exc}") from exc


def validate_or_repair_existing_line(
    client: Any,
    line_path: str,
    line: dict[str, Any],
    row: dict[str, Any],
    location_id: str,
) -> None:
    quantity = parse_decimal(line.get("quantity"), f"line {row['number']} quantity")
    if quantity != row["quantity"]:
        raise RuntimeError(f"Existing quantity for {row['number']} is {quantity}, expected {row['quantity']}.")
    price = parse_decimal(line.get("unitPrice"), f"line {row['number']} price")
    if abs(price - row["unitPrice"]) <= PRICE_TOLERANCE and clean_text(line.get("locationId")) == location_id:
        return
    client.patch(
        f"{line_path}({line['id']})",
        {"quantity": float(row["quantity"]), "unitPrice": float(row["unitPrice"]), "locationId": location_id},
    )
    refreshed = client.get_all(line_path, query={"$top": 1000})
    matches = [
        candidate for candidate in refreshed
        if clean_identifier(candidate.get("lineObjectNumber")) == row["number"]
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Could not verify repaired line {row['number']}.")
    validate_line_values(matches[0], row, location_id)


def validate_line_values(line: dict[str, Any], row: dict[str, Any], location_id: str) -> None:
    if clean_identifier(line.get("lineObjectNumber")) != row["number"]:
        raise RuntimeError(f"Line object mismatch for {row['number']}.")
    quantity = parse_decimal(line.get("quantity"), f"line {row['number']} quantity")
    price = parse_decimal(line.get("unitPrice"), f"line {row['number']} price")
    if quantity != row["quantity"]:
        raise RuntimeError(f"Line quantity for {row['number']} is {quantity}, expected {row['quantity']}.")
    if abs(price - row["unitPrice"]) > PRICE_TOLERANCE:
        raise RuntimeError(f"Line price for {row['number']} is {price}, expected {row['unitPrice']}.")
    if clean_text(line.get("locationId")) != location_id:
        raise RuntimeError(f"Line location for {row['number']} is not BOZ.")


def refresh_invoice_totals(client: Any, invoice_id: str, invoice_date: Any) -> None:
    if not clean_text(invoice_date):
        raise RuntimeError("Cannot refresh invoice totals because the invoice date is blank.")
    client.patch(
        client.company_path(f"salesInvoices({invoice_id})"),
        {"invoiceDate": invoice_date},
    )


def cell_value(values: tuple[Any, ...], index: int) -> Any:
    return values[index] if index < len(values) else None


def parse_decimal(value: Any, label: str) -> Decimal:
    if isinstance(value, Decimal):
        result = value
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        result = Decimal(str(value))
    else:
        text = clean_text(value).replace("\u00a0", "").replace(" ", "").replace("'", "")
        if "," in text and "." in text:
            text = text.replace(".", "").replace(",", ".") if text.rfind(",") > text.rfind(".") else text.replace(",", "")
        elif "," in text:
            text = text.replace(",", ".")
        try:
            result = Decimal(text)
        except (InvalidOperation, ValueError) as exc:
            raise RuntimeError(f"Invalid {label}: {value!r}") from exc
    if not result.is_finite():
        raise RuntimeError(f"Invalid {label}: {value!r}")
    return result


def decimal_json(value: Decimal) -> str:
    return format(value, "f")


def clean_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def clean_identifier(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(value)


def escape_odata(value: str) -> str:
    return value.replace("'", "''")


if __name__ == "__main__":
    raise SystemExit(main())
