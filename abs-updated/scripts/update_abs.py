#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
from collections import Counter
from copy import copy
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from collections.abc import Iterable

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator


REQUIRED_HEADERS = [
    "Description",
    "Unit Volume",
    "Alcohol Percentage Code",
    "Quantity",
    "Unit Price Excl. VAT",
    "Line Discount %",
    "Unit Cost",
]
FORMULA_COLUMNS = {7, 10, 14, 15, 18, 19, 20}
SUPPORT_CONSTANT_COLUMNS = {9, 16, 17}
INPUT_COLUMNS = {
    1: "doc",
    2: "desc",
    3: "unit_volume",
    4: "alc",
    5: "qty",
    6: "price",
    8: "discount",
    13: "unit_cost",
}


def natural_key(path: str | Path) -> tuple[str, int, str]:
    stem = Path(path).stem
    match = re.fullmatch(r"(.*?)(?: \((\d+)\))?", stem)
    base = (match.group(1) if match else stem).casefold()
    number = int(match.group(2)) if match and match.group(2) is not None else -1
    return base, number, Path(path).suffix.casefold()


def document_sort_key(value: Any) -> tuple[int, tuple[tuple[int, Any], ...]]:
    text = normalize_text(value)
    if not text:
        return 1, ()
    parts: list[tuple[int, Any]] = []
    for part in re.split(r"(\d+)", text):
        if not part:
            continue
        parts.append((0, int(part)) if part.isdigit() else (1, part.casefold()))
    return 0, tuple(parts)


def normalize_header(value: Any) -> str:
    return normalize_text(value).casefold()


def normalize_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalize_discount(value: Any, number_format: str | None = None) -> Any:
    if value in (None, ""):
        return 0
    if isinstance(value, (int, float, Decimal)):
        if number_format and "%" in number_format:
            return value
        if abs(value) >= 1:
            return value / 100
    return value


def source_sheet(workbook: Any) -> Any:
    return workbook["Lines"] if "Lines" in workbook.sheetnames else workbook.active


def header_index(sheet: Any) -> dict[str, int]:
    headers = {
        normalize_header(sheet.cell(1, col).value): col
        for col in range(1, sheet.max_column + 1)
        if sheet.cell(1, col).value is not None
    }
    missing = [header for header in REQUIRED_HEADERS if normalize_header(header) not in headers]
    if missing:
        raise ValueError(f"{sheet.title!r} is missing required headers: {', '.join(missing)}")
    return headers


def read_lines(paths: Iterable[str | Path]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for source_index, path in enumerate(sorted((Path(p) for p in paths), key=natural_key)):
        workbook = load_workbook(path, data_only=False, read_only=False)
        sheet = source_sheet(workbook)
        headers = header_index(sheet)
        document_col = headers.get(normalize_header("Document No."))
        discount_col = headers[normalize_header("Line Discount %")]
        last_document = ""

        def value(row: int, header: str) -> Any:
            column = headers.get(normalize_header(header))
            return sheet.cell(row, column).value if column else None

        for row in range(2, sheet.max_row + 1):
            raw_document = sheet.cell(row, document_col).value if document_col else None
            fields = {
                "desc": value(row, "Description"),
                "unit_volume": value(row, "Unit Volume"),
                "alc": value(row, "Alcohol Percentage Code"),
                "qty": value(row, "Quantity"),
                "price": value(row, "Unit Price Excl. VAT"),
                "unit_cost": value(row, "Unit Cost"),
            }
            discount_cell = sheet.cell(row, discount_col)
            fields["discount"] = normalize_discount(
                discount_cell.value,
                discount_cell.number_format,
            )
            if not normalize_text(raw_document) and all(
                item in (None, "") or (isinstance(item, str) and not item.strip())
                for item in fields.values()
            ):
                continue
            document = normalize_text(raw_document) or last_document
            if not document:
                raise ValueError(
                    f"{path.name}!{sheet.title} row {row} has content but no document number "
                    "and no preceding document to inherit"
                )
            last_document = document
            records.append(
                {
                    "doc": document,
                    **fields,
                    "source_order": (source_index, row),
                }
            )
    records.sort(
        key=lambda record: (
            document_sort_key(record["doc"]),
            record["source_order"],
        )
    )
    return records


def _fingerprint_value(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        try:
            return format(Decimal(str(value)).normalize(), "f")
        except InvalidOperation:
            pass
    return normalize_text(value).casefold()


def record_fingerprint(record: dict[str, Any]) -> tuple[str, ...]:
    return tuple(
        _fingerprint_value(record.get(key))
        for key in ("doc", "desc", "unit_volume", "alc", "qty", "price", "discount", "unit_cost")
    )


def existing_fingerprints(target: str | Path, sheet_name: str) -> Counter[tuple[str, ...]]:
    workbook = load_workbook(target, data_only=False, read_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Target sheet {sheet_name!r} was not found")
    sheet = workbook[sheet_name]
    result: Counter[tuple[str, ...]] = Counter()
    last_document = ""
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not values:
            continue
        document = normalize_text(values[0] if len(values) > 0 else None) or last_document
        if document:
            last_document = document
        record = {
            "doc": document,
            "desc": values[1] if len(values) > 1 else None,
            "unit_volume": values[2] if len(values) > 2 else None,
            "alc": values[3] if len(values) > 3 else None,
            "qty": values[4] if len(values) > 4 else None,
            "price": values[5] if len(values) > 5 else None,
            "discount": values[7] if len(values) > 7 else None,
            "unit_cost": values[12] if len(values) > 12 else None,
        }
        if document and any(record[key] not in (None, "") for key in record if key != "doc"):
            result[record_fingerprint(record)] += 1
    return result


def dedupe_records(
    records: Iterable[dict[str, Any]],
    existing: Counter[tuple[str, ...]],
) -> tuple[list[dict[str, Any]], int]:
    remaining = existing.copy()
    pending: list[dict[str, Any]] = []
    skipped = 0
    for record in records:
        fingerprint = record_fingerprint(record)
        if remaining[fingerprint] > 0:
            remaining[fingerprint] -= 1
            skipped += 1
        else:
            pending.append(record)
    return pending, skipped


def row_is_complete(sheet: Any, row: int) -> bool:
    return all(sheet.cell(row, col).value not in (None, "") for col in (1, 2, 5, 6, 13))


def detect_last_full_row(sheet: Any) -> int:
    for row in range(sheet.max_row, 1, -1):
        if row_is_complete(sheet, row):
            return row
    raise ValueError(f"Could not detect a complete data row in sheet {sheet.title!r}")


def copy_template_row(sheet: Any, template_row: int, target_row: int) -> None:
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[template_row].height
    for col in range(1, sheet.max_column + 1):
        source = sheet.cell(template_row, col)
        target = sheet.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        target.value = None
        if col in FORMULA_COLUMNS and isinstance(source.value, str) and source.value.startswith("="):
            target.value = Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)
        elif col in SUPPORT_CONSTANT_COLUMNS:
            target.value = source.value


def update_workbook_records(
    target: str | Path,
    records: Iterable[dict[str, Any]],
    output: str | Path,
    sheet_name: str = "Abs Updated",
    last_full_row: int | None = None,
) -> tuple[Path, int, int, int]:
    target_path = Path(target).resolve()
    output_path = Path(output).resolve()
    if target_path == output_path:
        raise ValueError("Output must differ from target; promote a validated copy separately")
    pending, _ = dedupe_records(records, existing_fingerprints(target_path, sheet_name))
    if not pending:
        raise ValueError("All supplied rows already exist; no workbook was written")

    workbook = load_workbook(target_path, data_only=False)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Target sheet {sheet_name!r} was not found")
    sheet = workbook[sheet_name]
    template_row = last_full_row or detect_last_full_row(sheet)
    start_row = template_row + 1
    template_formulas = {
        col for col in FORMULA_COLUMNS
        if isinstance(sheet.cell(template_row, col).value, str)
        and sheet.cell(template_row, col).value.startswith("=")
    }

    for offset, record in enumerate(pending):
        row = start_row + offset
        copy_template_row(sheet, template_row, row)
        for col, key in INPUT_COLUMNS.items():
            sheet.cell(row, col).value = record.get(key)
        sheet.cell(row, 11).value = None
        sheet.cell(row, 12).value = None

    calculation = getattr(workbook, "calculation", None)
    if calculation is not None:
        calculation.fullCalcOnLoad = True
        calculation.forceFullCalc = True
        calculation.calcMode = "auto"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    _verify_saved_rows(output_path, sheet_name, start_row, pending, template_formulas)
    return output_path, len(pending), start_row, start_row + len(pending) - 1


def update_workbook(
    target: str | Path,
    sources: Iterable[str | Path],
    output: str | Path,
    sheet_name: str = "Abs Updated",
    last_full_row: int | None = None,
) -> tuple[Path, int, int, int]:
    return update_workbook_records(
        target,
        read_lines(sources),
        output,
        sheet_name=sheet_name,
        last_full_row=last_full_row,
    )


def _verify_saved_rows(
    output: Path,
    sheet_name: str,
    start_row: int,
    records: list[dict[str, Any]],
    formula_columns: set[int],
) -> None:
    workbook = load_workbook(output, data_only=False, read_only=False)
    sheet = workbook[sheet_name]
    for offset, expected in enumerate(records):
        row = start_row + offset
        actual = {key: sheet.cell(row, col).value for col, key in INPUT_COLUMNS.items()}
        if record_fingerprint(actual) != record_fingerprint(expected):
            raise RuntimeError(f"Saved row {row} does not match its source record")
        for col in formula_columns:
            value = sheet.cell(row, col).value
            if not isinstance(value, str) or not value.startswith("="):
                raise RuntimeError(f"Formula did not extend to {sheet.cell(row, col).coordinate}")


def promote_current(validated_output: Path, current_copy: Path) -> None:
    current_copy.parent.mkdir(parents=True, exist_ok=True)
    temporary = current_copy.with_name(f".{current_copy.name}.tmp")
    shutil.copy2(validated_output, temporary)
    temporary.replace(current_copy)


def main() -> int:
    parser = argparse.ArgumentParser(description="Append Business Central Lines exports to ABs updated.xlsx.")
    parser.add_argument("sources", nargs="+", help="Source Lines*.xlsx workbooks")
    parser.add_argument("--target", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--sheet", default="Abs Updated")
    parser.add_argument("--last-full-row", type=int)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--json", action="store_true", dest="json_output")
    parser.add_argument("--current-copy", type=Path)
    args = parser.parse_args()

    target = args.target.resolve()
    output = args.output.resolve()
    if target == output:
        parser.error("--output must differ from --target")
    records = read_lines(args.sources)
    pending, skipped = dedupe_records(records, existing_fingerprints(target, args.sheet))
    summary = {
        "mode": "dry_run" if args.dry_run else "write",
        "sourceRows": len(records),
        "newRows": len(pending),
        "duplicateRowsSkipped": skipped,
        "target": str(target),
        "output": str(output),
    }
    if args.json_output:
        print(json.dumps(summary, indent=2, ensure_ascii=False))
    else:
        for key, value in summary.items():
            print(f"{key}: {value}")
    if args.dry_run or not pending:
        return 0

    saved, count, start, end = update_workbook_records(
        target,
        pending,
        output,
        sheet_name=args.sheet,
        last_full_row=args.last_full_row,
    )
    if args.current_copy:
        promote_current(saved, args.current_copy.resolve())
    if not args.json_output:
        print(f"Saved {saved}")
        print(f"Appended {count} rows at {start}-{end}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
